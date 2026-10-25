import io
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook, load_workbook

from app.apps.crm.lead_management import DEFAULTS
from app.apps.crm.models import CRMActivity, CRMLead, CRMContact, CRMTimelineEvent
from app.models.notification import Notification
from app.models.user import User

ROOT = "/api/v1/crm/lead-management"


def create(client, headers, **values):
    response = client.post(
        ROOT + "/leads",
        headers=headers,
        json={"values": {"full_name": "Lead Person", **values}},
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_lead_stage_rules_qualification_history_and_conflict(
    client, db, superuser_headers
):
    import copy

    settings = copy.deepcopy(DEFAULTS)
    settings["stages"][2]["required"] = ["need", "budget", "decisionMaker", "timeline"]
    saved = client.put(ROOT + "/settings", headers=superuser_headers, json=settings)
    assert saved.status_code == 200, saved.text
    lead = create(client, superuser_headers, email="stage@example.com")
    endpoint = ROOT + f"/leads/{lead['id']}"
    blocked = client.patch(
        endpoint, headers=superuser_headers, json={"values": {"status": "Qualified"}}
    )
    assert blocked.status_code == 422
    assert "need" in blocked.json()["detail"]["missingFields"]
    success = client.patch(
        endpoint,
        headers=superuser_headers,
        json={
            "values": {"status": "Qualified"},
            "version": lead["version"],
            "details": {
                "qualification": {
                    "need": "CRM",
                    "budget": "10000",
                    "decisionMaker": "Owner",
                    "timeline": "October",
                }
            },
        },
    )
    assert success.status_code == 200, success.text
    conflict = client.patch(
        endpoint,
        headers=superuser_headers,
        json={"version": lead["version"], "values": {"notes": "Stale edit"}},
    )
    assert conflict.status_code == 409
    assert (
        client.patch(
            endpoint,
            headers=superuser_headers,
            json={"values": {"status": "Converted"}},
        ).status_code
        == 422
    )
    detail = client.get(endpoint, headers=superuser_headers).json()
    change = next(e for e in detail["timeline"] if e["event_type"] == "stage_changed")
    assert change["metadata_json"]["from"] == "New"
    assert change["metadata_json"]["durationHours"] >= 0
    assert change["actor_user_id"]


def test_duplicate_prevention_and_transfer_tenant_attributes(client, superuser_headers):
    create(client, superuser_headers, email="same@example.com")
    duplicate = client.post(
        ROOT + "/leads",
        headers=superuser_headers,
        json={"values": {"full_name": "Duplicate", "email": "SAME@example.com"}},
    )
    assert duplicate.status_code == 409
    for field in [
        "organization_id",
        "is_converted",
        "created_by_user_id",
        "converted_deal_id",
    ]:
        response = client.post(
            ROOT + "/leads",
            headers=superuser_headers,
            json={"values": {"full_name": "Forbidden", field: 2}},
        )
        assert response.status_code == 422


def test_sequences_response_stop_and_follow_up(client, db, superuser_headers):
    lead = create(client, superuser_headers)
    endpoint = ROOT + f"/leads/{lead['id']}"
    enrollment = client.post(
        endpoint + "/sequence",
        headers=superuser_headers,
        json={"name": "Standard follow-up"},
    )
    assert enrollment.status_code == 200, enrollment.text
    assert (
        db.query(CRMActivity).filter_by(lead_id=lead["id"], status="Pending").count()
        == 3
    )
    due = (datetime.now(timezone.utc) + timedelta(days=1)).isoformat()
    activity = client.post(
        endpoint + "/activity",
        headers=superuser_headers,
        json={
            "type": "call",
            "outcome": "Interested",
            "responded": True,
            "nextFollowUp": due,
        },
    )
    assert activity.status_code == 200, activity.text
    detail = client.get(endpoint, headers=superuser_headers).json()
    assert detail["last_contacted_at"]
    assert not detail["details"]["sequence"]["active"]
    assert (
        db.query(CRMActivity).filter_by(lead_id=lead["id"], status="Cancelled").count()
        == 3
    )
    assert (
        db.query(CRMActivity).filter_by(lead_id=lead["id"], status="Pending").count()
        == 1
    )


def test_bulk_atomic_rollback(client, db, superuser_headers):
    first = create(client, superuser_headers)
    second = create(client, superuser_headers, full_name="Second")
    second_row = db.get(CRMLead, second["id"])
    second_row.status = "Converted"
    second_row.is_converted = True
    db.commit()
    response = client.post(
        ROOT + "/bulk/apply",
        headers=superuser_headers,
        json={
            "ids": [first["id"], second["id"]],
            "patch": {"values": {"status": "Contacted"}},
        },
    )
    assert response.status_code == 409, response.text
    db.expire_all()
    assert db.get(CRMLead, first["id"]).status == "New"


def make_excel(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    result = io.BytesIO()
    book.save(result)
    return result.getvalue()


def test_excel_preview_import_duplicate_skips_and_safe_export(
    client, db, superuser_headers
):
    create(client, superuser_headers, email="existing@example.com")
    contents = make_excel(
        [
            ["full_name", "email", "company_name", "phone", "phone_country_code"],
            ["New Lead", "new@example.com", "=Danger Company", "9876543210", "+91"],
            ["Existing Lead", "existing@example.com", "Company", None, None],
            ["Repeated row", "new@example.com", None, None, None],
        ]
    )
    # Excel formulas are rejected; ordinary text beginning with '=' is preserved by explicit string cells.
    book = load_workbook(io.BytesIO(contents))
    book.active["C2"].data_type = "s"
    output = io.BytesIO()
    book.save(output)
    response = client.post(
        ROOT + "/excel/preview",
        headers=superuser_headers,
        files={"file": ("leads.xlsx", output.getvalue())},
    )
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["ready"] == 1 and len(preview["skipped"]) == 2
    assert db.query(CRMLead).count() == 1  # Preview has no writes.
    response = client.post(
        ROOT + "/excel/import",
        headers=superuser_headers,
        json={"rows": preview["rows"]},
    )
    assert response.status_code == 200, response.text
    assert response.json()["created"] == 1
    response = client.get(ROOT + "/excel/export", headers=superuser_headers)
    assert response.status_code == 200, response.text
    exported = load_workbook(io.BytesIO(response.content))
    assert exported.active.max_row == 3
    assert exported.active.freeze_panes == "A2"
    cell = next(
        c for row in exported.active for c in row if c.value == "=Danger Company"
    )
    assert cell.data_type == "s"
    assert "9876543210" in [c.value for row in exported.active for c in row]
    assert (
        client.post(
            ROOT + "/excel/preview",
            headers=superuser_headers,
            files={"file": ("roundtrip.xlsx", response.content)},
        ).status_code
        == 200
    )


def test_excel_rejects_formulas_bad_headers_and_invalid_rows_atomically(
    client, db, superuser_headers
):
    for contents in [
        make_excel([["full_name", "email"], ["=1+1", "valid@example.com"]]),
        make_excel([["full_name", "organization_id"], ["Wrong", 2]]),
    ]:
        response = client.post(
            ROOT + "/excel/preview",
            headers=superuser_headers,
            files={"file": ("leads.xlsx", contents)},
        )
        assert response.status_code == 422
    response = client.post(
        ROOT + "/excel/import",
        headers=superuser_headers,
        json={
            "rows": [
                {"full_name": "Valid", "email": "good@example.com"},
                {"full_name": "Invalid", "email": "bad-email"},
            ]
        },
    )
    assert response.status_code == 422
    assert db.query(CRMLead).count() == 0


def test_sla_queue_aging_and_alert_idempotence(client, db, superuser_headers):
    lead = create(client, superuser_headers)
    row = db.get(CRMLead, lead["id"])
    row.created_at = datetime.now(timezone.utc) - timedelta(days=10)
    row.next_follow_up_at = datetime.now(timezone.utc) - timedelta(days=1)
    db.commit()
    response = client.get(ROOT + "/leads?queue=my_day", headers=superuser_headers)
    assert response.status_code == 200
    assert response.json()["items"][0]["signals"]["slaBreached"]
    assert response.json()["summary"]["stale"] == 1
    assert (
        client.get(ROOT + "/leads?queue=sla", headers=superuser_headers).json()["total"]
        == 1
    )
    assert (
        client.get(ROOT + "/leads?queue=active", headers=superuser_headers).json()[
            "total"
        ]
        == 1
    )
    one = client.post(ROOT + "/alerts/evaluate", headers=superuser_headers)
    two = client.post(ROOT + "/alerts/evaluate", headers=superuser_headers)
    assert one.json()["notificationsCreated"] == 1
    assert two.json()["notificationsCreated"] == 0
    assert db.query(Notification).filter_by(event_type="lead_sla").count() == 1


def test_conversion_reuses_contact_and_blocks_unqualified(
    client, db, superuser_headers
):
    lead = create(
        client, superuser_headers, email="convert@example.com", company_name="Company"
    )
    endpoint = ROOT + f"/leads/{lead['id']}"
    assert (
        client.post(
            endpoint + "/convert", headers=superuser_headers, json={}
        ).status_code
        == 422
    )
    contact = CRMContact(
        organization_id=1,
        first_name="Existing",
        full_name="Existing Contact",
        email="convert@example.com",
    )
    db.add(contact)
    db.commit()
    assert (
        client.patch(
            endpoint,
            headers=superuser_headers,
            json={"values": {"status": "Qualified"}},
        ).status_code
        == 200
    )
    preview = client.get(endpoint + "/conversion-preview", headers=superuser_headers)
    assert preview.json()["matches"]["contacts"][0]["id"] == contact.id
    assert (
        client.post(
            endpoint + "/convert", headers=superuser_headers, json={}
        ).status_code
        == 409
    )
    response = client.post(
        endpoint + "/convert",
        headers=superuser_headers,
        json={"existingContactId": contact.id, "createDeal": False},
    )
    assert response.status_code == 200, response.text
    assert response.json()["contact"]["id"] == contact.id
    assert db.query(CRMContact).count() == 1


def test_assignment_rules_and_workflows(client, db, superuser_headers):
    import copy

    settings = copy.deepcopy(DEFAULTS)
    user = db.query(User).filter_by(is_superuser=True).first()
    settings["assignment"] = {"mode": "round_robin", "ownerIds": [user.id], "rules": []}
    settings["workflows"] = [
        {
            "name": "New lead task",
            "trigger": "created",
            "action": "task",
            "title": "Prepare discovery call",
            "delayHours": 2,
        }
    ]
    response = client.put(ROOT + "/settings", headers=superuser_headers, json=settings)
    assert response.status_code == 200, response.text
    lead = create(client, superuser_headers)
    assert lead["owner_user_id"] == user.id
    assert (
        db.query(CRMActivity)
        .filter_by(
            lead_id=lead["id"], subject="Prepare discovery call", status="Pending"
        )
        .count()
        == 1
    )


def test_archive_and_reactivation(client, superuser_headers):
    lead = create(client, superuser_headers)
    endpoint = ROOT + f"/leads/{lead['id']}"
    assert (
        client.patch(
            endpoint, headers=superuser_headers, json={"details": {"archived": True}}
        ).status_code
        == 422
    )
    response = client.patch(
        endpoint,
        headers=superuser_headers,
        json={
            "details": {
                "archived": True,
                "archiveReason": "Review later",
                "reviewAt": "2026-01-01T00:00:00Z",
            }
        },
    )
    assert response.status_code == 200, response.text
    assert client.get(ROOT + "/leads", headers=superuser_headers).json()["total"] == 0
    assert (
        client.get(ROOT + "/leads?queue=archived", headers=superuser_headers).json()[
            "total"
        ]
        == 1
    )
    assert (
        client.get(
            ROOT + "/leads?queue=reactivation", headers=superuser_headers
        ).json()["total"]
        == 1
    )


def test_existing_lead_zero_version_edit_and_manual_score_reason(
    client, db, superuser_headers
):
    row = CRMLead(
        organization_id=1, first_name="Legacy", full_name="Legacy Lead", status="New"
    )
    db.add(row)
    db.commit()
    endpoint = ROOT + f"/leads/{row.id}"
    response = client.patch(
        endpoint,
        headers=superuser_headers,
        json={"version": 0, "values": {"notes": "New note"}},
    )
    assert response.status_code == 200, response.text
    assert (
        client.patch(
            endpoint, headers=superuser_headers, json={"values": {"lead_score": 91}}
        ).status_code
        == 422
    )
    response = client.patch(
        endpoint,
        headers=superuser_headers,
        json={
            "values": {"lead_score": 91},
            "details": {"scoreOverrideReason": "Confirmed buying intent"},
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["lead_score_mode"] == "manual"


def test_lead_management_rbac_and_export_visibility(client, db, superuser_headers):
    from tests.test_crm_phase_1_core_foundation import _headers_for_permissions

    headers = _headers_for_permissions(
        client, db, "scope-user@example.com", ["crm_leads_manage"]
    )
    hidden = create(client, superuser_headers, email="hidden@example.com")
    own = create(client, headers, email="own@example.com")
    foreign = CRMLead(
        organization_id=99,
        first_name="Foreign",
        full_name="Foreign Lead",
        email="foreign@example.com",
    )
    db.add(foreign)
    db.commit()
    listing = client.get(ROOT + "/leads", headers=headers)
    assert [i["id"] for i in listing.json()["items"]] == [own["id"]]
    for ident in [hidden["id"], foreign.id]:
        assert client.get(ROOT + f"/leads/{ident}", headers=headers).status_code == 404
        assert (
            client.patch(
                ROOT + f"/leads/{ident}",
                headers=headers,
                json={"values": {"notes": "No"}},
            ).status_code
            == 404
        )
    response = client.get(ROOT + "/excel/export", headers=headers)
    assert response.status_code == 200
    values = [
        cell.value
        for row in load_workbook(io.BytesIO(response.content)).active
        for cell in row
    ]
    assert (
        "own@example.com" in values
        and "hidden@example.com" not in values
        and "foreign@example.com" not in values
    )
    assert (
        client.put(ROOT + "/settings", headers=headers, json=DEFAULTS).status_code
        == 403
    )
    reader = _headers_for_permissions(
        client, db, "scope-reader@example.com", ["crm_leads_view"]
    )
    assert client.get(ROOT + "/settings", headers=reader).json()["canManage"] is False
    assert (
        client.post(
            ROOT + "/leads", headers=reader, json={"values": {"full_name": "No"}}
        ).status_code
        == 403
    )


def test_website_capture_is_tenant_owned_consent_gated_and_revocable(
    client, db, superuser_headers
):
    user = db.query(User).filter_by(is_superuser=True).first()
    enabled = client.put(
        ROOT + "/capture-settings",
        headers=superuser_headers,
        json={"enabled": True, "ownerId": user.id},
    )
    assert enabled.status_code == 200, enabled.text
    path = enabled.json()["path"]
    assert client.get(path).status_code == 200
    assert (
        client.post(
            path,
            data={
                "full_name": "Website Lead",
                "email": "web@example.com",
                "message": "Need a demo",
            },
        ).status_code
        == 422
    )
    values = {
        "full_name": "Website Lead",
        "email": "web@example.com",
        "message": "Need a demo",
        "consent": "yes",
    }
    assert client.post(path, data=values).status_code == 200
    lead = db.query(CRMLead).filter_by(email="web@example.com").one()
    assert lead.organization_id == 1 and lead.owner_user_id == user.id
    inbox = client.get(
        ROOT + "/leads?queue=unassigned", headers=superuser_headers
    ).json()
    assert inbox["items"][0]["id"] == lead.id
    assert inbox["items"][0]["details"]["contactConsentAt"]
    assert client.post(path, data=values).status_code == 200
    assert db.query(CRMLead).filter_by(email="web@example.com").count() == 1
    assert (
        "capture"
        not in client.get(ROOT + "/settings", headers=superuser_headers).json()[
            "settings"
        ]
    )
    client.put(
        ROOT + "/capture-settings",
        headers=superuser_headers,
        json={"enabled": False, "ownerId": user.id},
    )
    assert client.get(path).status_code == 404


def test_booking_links_single_use_and_calendar_conflicts(client, db, superuser_headers):
    from app.apps.crm.models import CRMMeeting

    lead = create(client, superuser_headers)
    path = client.post(
        ROOT + f"/leads/{lead['id']}/booking-link", headers=superuser_headers
    ).json()["path"]
    assert client.get(path).status_code == 200
    start = (datetime.now(timezone.utc) + timedelta(days=1)).replace(
        second=0, microsecond=0
    )
    data = {"start_time": start.strftime("%Y-%m-%dT%H:%M"), "offset": "+00:00"}
    booked = client.post(path, data=data)
    assert booked.status_code == 200, booked.text
    assert db.query(CRMMeeting).filter_by(lead_id=lead["id"]).count() == 1
    assert client.post(path, data=data).status_code == 404
    second = create(client, superuser_headers, full_name="Second Lead")
    other = client.post(
        ROOT + f"/leads/{second['id']}/booking-link", headers=superuser_headers
    ).json()["path"]
    assert client.post(other, data=data).status_code == 409
    assert db.query(CRMMeeting).count() == 1


def test_analytics_uses_recorded_first_response_and_visible_stage_history(
    client, superuser_headers
):
    lead = create(client, superuser_headers, source="Referral")
    endpoint = ROOT + f"/leads/{lead['id']}"
    client.post(
        endpoint + "/activity",
        headers=superuser_headers,
        json={"type": "call", "outcome": "Qualified discussion"},
    )
    client.patch(
        endpoint, headers=superuser_headers, json={"values": {"status": "Qualified"}}
    )
    response = client.get(ROOT + "/analytics", headers=superuser_headers)
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["bySource"][0]["responseSampleSize"] == 1
    assert data["bySource"][0]["qualified"] == 1
    assert (
        next(s for s in data["stageHistory"] if s["stage"] == "New")["timedExits"] == 1
    )


def test_classic_api_cannot_bypass_configured_stage_requirements(
    client, superuser_headers
):
    import copy

    settings = copy.deepcopy(DEFAULTS)
    settings["stages"][2]["required"] = ["budget"]
    assert (
        client.put(
            ROOT + "/settings", headers=superuser_headers, json=settings
        ).status_code
        == 200
    )
    lead = create(client, superuser_headers)
    response = client.patch(
        f"/api/v1/crm/leads/{lead['id']}",
        headers=superuser_headers,
        json={"status": "Qualified"},
    )
    assert response.status_code == 422
    assert (
        client.post(
            f"/api/v1/crm/leads/{lead['id']}/qualify", headers=superuser_headers
        ).status_code
        == 422
    )


def test_scheduled_sla_alerts_are_bounded_idempotent_and_skip_archived(
    client, db, superuser_headers
):
    from app.apps.crm.lead_management import scheduled_lead_alerts

    lead = create(client, superuser_headers)
    archived = create(client, superuser_headers, full_name="Archived Lead")
    endpoint = ROOT + f"/leads/{archived['id']}"
    assert (
        client.patch(
            endpoint,
            headers=superuser_headers,
            json={"details": {"archived": True, "archiveReason": "No longer active"}},
        ).status_code
        == 200
    )
    for ident in [lead["id"], archived["id"]]:
        db.get(CRMLead, ident).created_at = datetime.now(timezone.utc) - timedelta(
            days=2
        )
    db.commit()
    result = scheduled_lead_alerts(db)
    assert result["leadsEscalated"] == 1
    assert scheduled_lead_alerts(db)["leadsEscalated"] == 0


def test_nulls_and_invalid_score_return_validation_errors(client, superuser_headers):
    lead = create(client, superuser_headers)
    endpoint = ROOT + f"/leads/{lead['id']}"
    for field in ["full_name", "first_name", "status", "lead_score"]:
        response = client.patch(
            endpoint, headers=superuser_headers, json={"values": {field: None}}
        )
        assert response.status_code == 422, response.text


def test_sequence_task_completion_advances_follow_up_and_finishes_sequence(
    client, db, superuser_headers
):
    lead = create(client, superuser_headers)
    endpoint = ROOT + f"/leads/{lead['id']}"
    client.post(
        endpoint + "/sequence",
        headers=superuser_headers,
        json={"name": "Standard follow-up"},
    )
    tasks = (
        db.query(CRMActivity)
        .filter_by(lead_id=lead["id"], status="Pending")
        .order_by(CRMActivity.due_date)
        .all()
    )
    for task in tasks:
        response = client.post(
            endpoint + f"/activities/{task.id}/complete", headers=superuser_headers
        )
        assert response.status_code == 200, response.text
    response = client.get(endpoint, headers=superuser_headers).json()
    assert response["details"]["sequence"]["active"] is False
    assert response["next_follow_up_at"] is None
    assert (
        client.post(
            endpoint + f"/activities/{tasks[0].id}/complete", headers=superuser_headers
        ).status_code
        == 409
    )


def test_excel_preserves_qualification_loss_and_archive_details(
    client, superuser_headers
):
    response = client.post(
        ROOT + "/excel/import",
        headers=superuser_headers,
        json={
            "rows": [
                {
                    "full_name": "Lost Prospect",
                    "email": "lost-import@example.com",
                    "status": "Lost",
                    "loss_reason": "Budget",
                    "need": "Business automation",
                    "budget": "100000",
                    "decision_maker": "Director",
                    "purchase_timeline": "Next quarter",
                    "archived": True,
                    "archive_reason": "Revisit when funding arrives",
                    "review_at": "2026-12-01T09:00:00Z",
                }
            ]
        },
    )
    assert response.status_code == 200, response.text
    ident = response.json()["ids"][0]
    detail = client.get(ROOT + f"/leads/{ident}", headers=superuser_headers).json()
    assert detail["details"]["qualification"]["decisionMaker"] == "Director"
    assert detail["details"]["lossReason"] == "Budget"
    response = client.get(ROOT + "/excel/export", headers=superuser_headers)
    sheet = load_workbook(io.BytesIO(response.content)).active
    headers = [cell.value for cell in sheet[1]]
    values = dict(zip(headers, [cell.value for cell in sheet[2]]))
    assert values["decision_maker"] == "Director" and values["loss_reason"] == "Budget"
    assert values["archived"] is True


def test_product_routing_and_duplicate_numbers_with_different_calling_codes(
    client, db, superuser_headers
):
    import copy

    settings = copy.deepcopy(DEFAULTS)
    owner = db.query(User).filter_by(is_superuser=True).first()
    settings["assignment"]["rules"] = [
        {"field": "product", "equals": "Implementation", "ownerId": owner.id}
    ]
    assert (
        client.put(
            ROOT + "/settings", headers=superuser_headers, json=settings
        ).status_code
        == 200
    )
    response = client.post(
        ROOT + "/leads",
        headers=superuser_headers,
        json={
            "values": {
                "full_name": "Product Lead",
                "phone": "9876543210",
                "phone_country_code": "+91",
            },
            "details": {
                "interests": [
                    {
                        "name": "Implementation",
                        "value": 20000,
                        "requirements": "Training",
                    }
                ]
            },
        },
    )
    assert response.status_code == 201, response.text
    assert response.json()["owner_user_id"] == owner.id
    response = client.post(
        ROOT + "/leads",
        headers=superuser_headers,
        json={
            "values": {
                "full_name": "Different Country",
                "phone": "9876543210",
                "phone_country_code": "+1",
            }
        },
    )
    assert response.status_code == 201, response.text
