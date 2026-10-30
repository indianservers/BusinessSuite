from datetime import date
from io import BytesIO
import pytest
from openpyxl import load_workbook
from app.core.security import create_access_token
from app.models.user import User, Role, Permission
from app.models.employee import Employee, EmployeeEducation
from app.models.audit import AuditLog
from app.apps.interns_hub.models import InternApplication
from tests.test_interns_hub import interns

EMP = "/api/v1/employees/export/full"
INTERN = "/api/v1/interns-hub/management/export/full"
DIRECTORY = "/api/v1/employees/directory/export/full"

def workbook(response):
    assert response.status_code == 200, response.text
    assert "no-store" in response.headers["cache-control"]
    return load_workbook(BytesIO(response.content), data_only=False)

@pytest.mark.parametrize("path", [EMP, INTERN, DIRECTORY])
def test_full_export_rejects_non_superadmin_even_with_export_permission(client, db, path):
    role = Role(name="super_admin")
    role.permissions = [Permission(name="employee_export"), Permission(name="employee_view")]
    user = User(email="delegated-export@example.com", hashed_password="", is_active=True, role=role, is_superuser=False)
    db.add(user); db.commit()
    headers = {"Authorization": "Bearer " + create_access_token(user.id)}
    assert client.get(path, headers=headers).status_code == 403
    assert client.get(path).status_code in {401, 403}
    assert db.query(AuditLog).filter_by(action="EXPORT").count() == 0


def test_full_export_includes_profiles_children_and_all_pages(client, db, superuser_headers):
    rows = [Employee(employee_id=f"FULL-{i:03}", first_name="Export", last_name=str(i), date_of_joining=date(2026, 1, 1),
        personal_email=f"person{i}@example.com", account_number="001234567890", aadhaar_number="123456789012",
        phone_number="+919876543210", employment_type="Full-time") for i in range(205)]
    db.add_all(rows); db.flush()
    rows[0].bio = '=HYPERLINK("https://example.com")'
    db.add(EmployeeEducation(employee_id=rows[0].id, degree="B.Tech", institution="Export College"))
    db.add(Employee(employee_id="DELETED-FULL", first_name="Deleted", last_name="Person", date_of_joining=date(2026, 1, 1), deleted_at=__import__("datetime").datetime.now()))
    db.commit()
    response = client.get(EMP, headers=superuser_headers, params={"search": "FULL-", "page": 2})
    book = workbook(response)
    assert response.headers["x-total-records"] == "205"
    sheet = book["Profiles"]
    data = list(sheet.values)
    assert len(data) == 206
    columns = dict(zip(data[0], data[1]))
    assert columns["account_number"] == "001234567890"
    assert columns["aadhaar_number"] == "123456789012"
    assert columns["phone_number"] == "+919876543210"
    cell = sheet.cell(2, data[0].index("bio") + 1)
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
    assert book["Education"].max_row == 2
    assert "Salary" in book.sheetnames and "Family" in book.sheetnames
    assert db.query(AuditLog).filter_by(action="EXPORT", endpoint=EMP).count() == 1
    filtered = workbook(client.get(EMP, headers=superuser_headers, params={"search": "FULL-003"}))
    assert filtered["Profiles"].max_row == 2


def test_full_intern_export_preserves_filters_and_excludes_auth_secrets(client, db, interns, superuser_headers):
    person = interns[0][0]
    db.add(InternApplication(reference="EXPORT-APP", email="application@example.com", full_name="Intern Name",
        password_hash="MUST-NOT-EXPORT-PASSWORD", details={"college": "Example University"},
        identity_number_encrypted="MUST-NOT-EXPORT-IDENTITY", identity_last4="1234", employee_id=person.id))
    db.commit()
    response = client.get(INTERN, headers=superuser_headers, params={"search": person.employee_id})
    book = workbook(response)
    assert book["Profiles"].max_row == 2
    assert book["Applications"].max_row == 2
    for sheet in book:
        for row in sheet.values:
            assert "password_hash" not in row
            assert "identity_number_encrypted" not in row
            assert all("MUST-NOT-EXPORT" not in str(value) for value in row)
    assert book["Accounts"].max_row == 2
    empty = workbook(client.get(INTERN, headers=superuser_headers, params={"search": "NO-SUCH-PERSON"}))
    assert empty["Profiles"].max_row == 1


def test_intern_export_keeps_organization_boundary(client, db, interns, superuser_headers):
    from app.models.company import Company, Branch
    company = Company(name="Other organization")
    db.add(company); db.flush()
    branch = Branch(name="Other branch", company_id=company.id)
    db.add(branch); db.flush()
    db.add(Employee(employee_id="OTHER-TENANT", first_name="Other", last_name="Intern",
        employment_type="Intern", branch_id=branch.id, date_of_joining=date(2026, 1, 1)))
    db.commit()
    book = workbook(client.get(INTERN, headers=superuser_headers))
    codes = [row[1] for row in list(book["Profiles"].values)[1:]]
    assert "OTHER-TENANT" not in codes
    assert len(codes) == 2


def test_directory_full_export_filters_and_keeps_long_text(client, db, superuser_headers):
    bio = "Long profile text " * 3000
    db.add_all([
        Employee(employee_id="DIR-FULL-A", first_name="Directory", last_name="Match", skills_tags="Python", bio=bio, date_of_joining=date(2026, 1, 1)),
        Employee(employee_id="DIR-FULL-B", first_name="Directory", last_name="Other", skills_tags="Java", date_of_joining=date(2026, 1, 1)),
    ])
    db.commit()
    book = workbook(client.get(DIRECTORY, headers=superuser_headers, params={"skills": "Python"}))
    assert book["Profiles"].max_row == 2
    chunks = list(book["Long Text"].values)[1:]
    assert "".join(row[4] for row in chunks) == bio
