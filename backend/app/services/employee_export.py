"""Super Admin HRMS profile exports. Never serialize authentication records."""
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
import json

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import joinedload

from app.models.audit import AuditLog
from app.models.user import User, Role
from app.models.employee import (
    Employee, EmployeeEducation, EmployeeExperience, EmployeeSkill,
    EmployeeCertification, EmployeeFamilyMember, EmployeeDocument,
    EmployeeLifecycleEvent,
)
from app.models.document import EmployeeCertificate
from app.models.payroll import EmployeeSalary
from app.models.platform import CustomFieldDefinition, CustomFieldValue
from app.apps.interns_hub.models import InternApplication
from app.apps.interns_hub.management_models import InternBatch, InternBatchMembership


def full_profile_workbook(db, actor, query, *, filename, endpoint):
    # Defense in depth: callers cannot accidentally expose this through a weaker route.
    from fastapi import HTTPException
    if not actor.is_active or not actor.is_superuser:
        raise HTTPException(403, "Super Admin access required for full-details export")
    selected = query.with_entities(Employee.id).order_by(None).distinct().subquery()
    ids = select(selected.c.id)
    book = Workbook(write_only=True)
    long_text = []

    def sheet(title, headers, rows):
        ws = book.create_sheet(title)
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        def cells(values, row_number):
            result = []
            for column, value in enumerate(values, 1):
                if isinstance(value, (date, datetime)):
                    value = value.isoformat()
                elif isinstance(value, Decimal):
                    value = str(value)
                elif isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False, default=str)
                if isinstance(value, str):
                    value = ILLEGAL_CHARACTERS_RE.sub("", value)
                    if len(value) > 32767:
                        parts = [value[i:i + 32000] for i in range(0, len(value), 32000)]
                        long_text.extend((title, row_number, headers[column - 1], n, part)
                                         for n, part in enumerate(parts, 1))
                        value = f"See Long Text sheet ({len(parts)} parts)"
                cell = WriteOnlyCell(ws, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Prevent spreadsheet formulas in user-entered data.
                if row_number == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="173B56")
                result.append(cell)
            return result
        ws.append(cells(headers, 1))
        count = 0
        for count, values in enumerate(rows, 1):
            ws.append(cells(values, count + 1))
        from openpyxl.utils import get_column_letter
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count + 1}"
        return count

    columns = [column.name for column in Employee.__table__.columns]
    people = {}
    def profiles():
        for employee in db.query(Employee).filter(Employee.id.in_(ids)).options(
            joinedload(Employee.branch), joinedload(Employee.department),
            joinedload(Employee.designation), joinedload(Employee.reporting_manager),
        ).order_by(Employee.employee_id, Employee.id).yield_per(500):
            name = f"{employee.first_name} {employee.last_name}".strip()
            people[employee.id] = (employee.employee_id, name)
            labels = [getattr(getattr(employee, key), "name", None)
                      for key in ("branch", "department", "designation")]
            manager = employee.reporting_manager
            labels.append(f"{manager.first_name} {manager.last_name}" if manager else None)
            yield [getattr(employee, key) for key in columns] + labels
    total = sheet("Profiles", columns + ["branch_name", "department_name", "designation_name", "manager_name"], profiles())
    accounts = db.query(Employee.id, User.id, User.email, User.is_active, User.is_superuser, Role.name).join(
        User, User.id == Employee.user_id
    ).outerjoin(Role, Role.id == User.role_id).filter(Employee.id.in_(ids)).order_by(Employee.id)
    sheet("Accounts", ["employee_id", "user_id", "email", "active", "super_admin", "role"], accounts.yield_per(500))

    def related(title, model, excluded=()):
        fields = [column.name for column in model.__table__.columns if column.name not in excluded]
        rows = db.query(model).filter(model.employee_id.in_(ids)).order_by(model.employee_id, model.id).yield_per(500)
        sheet(title, ["employee_code", "employee_name"] + fields,
              (list(people[row.employee_id]) + [getattr(row, key) for key in fields] for row in rows))
    for title, model in (
        ("Education", EmployeeEducation), ("Experience", EmployeeExperience),
        ("Skills", EmployeeSkill), ("Qualifications", EmployeeCertification),
        ("Family", EmployeeFamilyMember), ("Documents", EmployeeDocument),
        ("Certificates", EmployeeCertificate), ("Lifecycle", EmployeeLifecycleEvent),
        ("Salary", EmployeeSalary), ("Batch Memberships", InternBatchMembership),
    ):
        related(title, model)
    related("Applications", InternApplication, excluded=("password_hash", "identity_number_encrypted"))
    memberships = db.query(InternBatchMembership.batch_id).filter(InternBatchMembership.employee_id.in_(ids))
    fields = [c.name for c in InternBatch.__table__.columns]
    sheet("Batches", fields, ([getattr(row, key) for key in fields] for row in
          db.query(InternBatch).filter(InternBatch.id.in_(memberships)).order_by(InternBatch.id).yield_per(500)))
    custom = db.query(CustomFieldValue, CustomFieldDefinition).join(
        CustomFieldDefinition, CustomFieldDefinition.id == CustomFieldValue.definition_id
    ).filter(CustomFieldValue.entity_type == "employee", CustomFieldValue.entity_id.in_(ids)).order_by(CustomFieldValue.id)
    sheet("Custom Fields", ["employee_id", "employee_code", "field_key", "label", "value_text", "value_json"],
          ([value.entity_id, people[value.entity_id][0], definition.field_key, definition.label,
            value.value_text, value.value_json] for value, definition in custom.yield_per(500)))
    if long_text:
        sheet("Long Text", ["sheet", "row", "column", "part", "text"], iter(long_text))
    sheet("Export Info", ["item", "value"], iter([
        ("Exported at (UTC)", datetime.now(timezone.utc).isoformat()),
        ("Exported by", actor.email), ("Profiles", total),
        ("Scope", "All matching non-deleted profiles; current list filters apply, without pagination."),
        ("Documents", "Document metadata and stored references are included; attachment binaries are not embedded."),
        ("Excluded", "Passwords, password hashes, login/reset tokens and encrypted application identity payloads."),
    ]))
    output = BytesIO()
    book.save(output)
    db.add(AuditLog(user_id=actor.id, method="GET", endpoint=endpoint, status_code=200,
                    entity_type="employee", action="EXPORT",
                    description=f"Super Admin exported {total} full HRMS profiles"))
    db.commit()
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"',
                             "Cache-Control": "no-store", "X-Total-Records": str(total)})
