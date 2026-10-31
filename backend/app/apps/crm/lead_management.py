"""Lead operations. All reads, mutations and transfers share CRM tenant/RBAC rules."""

from __future__ import annotations
from app.core.rich_text import rich_text_type

import copy
import csv
import io
import html
import secrets
import re
import zipfile
from collections import Counter
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from typing import Any, Literal

from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    Query,
    UploadFile,
)
from fastapi.responses import Response, HTMLResponse
from pydantic import BaseModel, ConfigDict, Field, field_validator
from sqlalchemy import and_, func, or_
from sqlalchemy.orm import Session

from app.core.deps import get_db, RequirePermission
from app.models.user import User
from app.models.notification import Notification
from app.apps.crm.models import (
    CRMLead,
    CRMActivity,
    CRMContact,
    CRMCompany,
    CRMMeeting,
    CRMTimelineEvent,
)
from app.apps.crm.lead_models import CRMLeadManagementConfig, CRMLeadOperatingState

router = APIRouter(prefix="/lead-management", tags=["Lead management"])
ACCESS = ("crm_view", "crm_manage", "crm_admin", "crm_leads_view", "crm_leads_manage")
WRITE = ("crm_manage", "crm_admin", "crm_leads_manage")
MAX_ROWS = 5000
MAX_FILE = 3 * 1024 * 1024
BASE_TRANSFER_FIELDS = [
    "full_name",
    "email",
    "phone",
    "phone_country_code",
    "company_name",
    "job_title",
    "source",
    "status",
    "estimated_value",
    "next_follow_up_at",
    "city",
    "state",
    "country",
    "industry",
    "notes",
]
DETAIL_TRANSFER_FIELDS = [
    "need",
    "budget",
    "decision_maker",
    "purchase_timeline",
    "loss_reason",
    "competitor",
    "archived",
    "archive_reason",
    "review_at",
]
TRANSFER_FIELDS = BASE_TRANSFER_FIELDS + DETAIL_TRANSFER_FIELDS
DEFAULTS = {
    "stages": [
        {
            "name": "New",
            "required": [],
            "allowedNext": ["Contacted", "Qualified", "Lost"],
        },
        {"name": "Contacted", "required": [], "allowedNext": ["Qualified", "Lost"]},
        {"name": "Qualified", "required": [], "allowedNext": ["Lost"]},
        {"name": "Converted", "required": [], "allowedNext": []},
        {"name": "Lost", "required": ["lossReason"], "allowedNext": ["New"]},
    ],
    "slaHours": 24,
    "staleDays": 7,
    "preventDuplicates": True,
    "assignment": {"mode": "manual", "ownerIds": [], "rules": []},
    "sequences": [
        {
            "name": "Standard follow-up",
            "steps": [
                {
                    "day": 0,
                    "type": "call",
                    "title": "Introduce yourself and qualify the enquiry",
                },
                {"day": 2, "type": "email", "title": "Send a tailored follow-up"},
                {"day": 5, "type": "call", "title": "Confirm next steps"},
            ],
        }
    ],
    "workflows": [],
}


def core():
    # Deferred import keeps the existing generic CRM router independent.
    from app.apps.crm.api import router as crm

    return crm


def now():
    return datetime.now(timezone.utc)


def dt(value):
    if not value:
        return None
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            raise HTTPException(422, "Invalid date/time")
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value


def config(db, org, lock=False):
    query = db.query(CRMLeadManagementConfig).filter_by(organization_id=org)
    row = (query.with_for_update() if lock else query).first()
    return {
        **copy.deepcopy(DEFAULTS),
        **copy.deepcopy(row.settings_json if row else {}),
    }


def public_page(title, content):
    return HTMLResponse(
        '<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
        "<title>"
        + html.escape(title)
        + "</title><style>body{font:16px system-ui;color:#182230;background:#f8fafc;margin:0;padding:24px}main{max-width:480px;margin:30px auto;background:white;padding:28px;border:1px solid #e2e8f0;border-radius:12px}h1{font-size:24px}label{display:block;margin:18px 0;font-size:14px}input,textarea,select,button{box-sizing:border-box;width:100%;padding:12px;margin-top:7px;border:1px solid #cbd5e1;border-radius:6px;font:inherit}button{background:#182230;color:white;cursor:pointer}input[type=checkbox]{width:auto}small,p{color:#475569;line-height:1.6}.trap{position:absolute;left:-9999px}</style><main><h1>"
        + html.escape(title)
        + "</h1>"
        + content
        + "</main></html>",
        headers={
            "Cache-Control": "no-store",
            "Referrer-Policy": "no-referrer",
            "X-Content-Type-Options": "nosniff",
            "Content-Security-Policy": "default-src 'none'; style-src 'unsafe-inline'; form-action 'self'; frame-ancestors 'self'",
        },
    )


def visible_query(db, user):
    crm = core()
    crm._assert_entity_permission("leads", user, "view")
    resource = crm._get_resource("leads")
    return crm._apply_record_visibility(
        crm._base_query(db, resource, crm._organization_id(db, user)),
        resource,
        user,
        db,
    )


def lead_record(db, user, ident, write=False, lock=False):
    if write:
        core()._assert_entity_permission("leads", user, "manage")
    query = visible_query(db, user).filter(CRMLead.id == ident)
    lead = (query.with_for_update() if lock else query).first()
    if not lead:
        raise HTTPException(404, "Lead was not found in your permitted scope")
    return lead


def state(db, lead, create=True):
    row = (
        db.query(CRMLeadOperatingState)
        .filter_by(organization_id=lead.organization_id, lead_id=lead.id)
        .first()
    )
    if not row and create:
        row = CRMLeadOperatingState(
            organization_id=lead.organization_id,
            lead_id=lead.id,
            data_json={},
            version=1,
        )
        db.add(row)
        db.flush()
    return row


def event(db, lead, user, kind, title, metadata=None, body=None):
    core()._create_timeline_activity(
        db,
        organization_id=lead.organization_id,
        entity_type="lead",
        entity_id=lead.id,
        activity_type=kind,
        title=title,
        body=body,
        user_id=user.id,
        metadata=metadata or {},
    )


def check_owner(db, user, owner_id):
    candidate = db.get(User, owner_id)
    crm = core()
    from app.apps.crm.lead_assignment import validate_assignee

    validate_assignee(db, candidate, crm._organization_id(db, user))
    crm._assert_write_scope(
        crm._get_resource("leads"), {"owner_user_id": owner_id}, user, db
    )
    return candidate


def duplicates(db, org, values, exclude=None):
    email = str(values.get("email") or "").strip().lower()
    phone = re.sub(r"\D", "", str(values.get("phone") or ""))
    clauses = []
    if email:
        clauses.append(func.lower(CRMLead.email) == email)
    # Normalize legacy punctuation as well as new phone values.
    if phone:
        normalized = CRMLead.phone
        for char in ["+", " ", "-", "(", ")", "."]:
            normalized = func.replace(normalized, char, "")
        phone_match = normalized == phone
        code = values.get("phone_country_code")
        if code:
            phone_match = and_(
                phone_match,
                or_(
                    CRMLead.phone_country_code == code,
                    CRMLead.phone_country_code.is_(None),
                ),
            )
        clauses.append(phone_match)
    if not clauses:
        return []
    query = db.query(CRMLead.id).filter(
        CRMLead.organization_id == org, CRMLead.deleted_at.is_(None), or_(*clauses)
    )
    if exclude:
        query = query.filter(CRMLead.id != exclude)
    return [r.id for r in query.limit(10)]


def validate_transition(settings, lead, values, details, converting=False):
    target = values.get("status", lead.status if lead else "New") or "New"
    stages = {s["name"]: s for s in settings["stages"]}
    if target not in stages:
        raise HTTPException(422, "Choose a configured lead stage")
    if target == "Converted" and not converting:
        if not lead or not lead.is_converted:
            raise HTTPException(
                422, "Use Convert Lead to create or link customer records"
            )
    if lead and lead.is_converted and target != "Converted":
        raise HTTPException(
            409, "A converted lead cannot return to the prospect pipeline"
        )
    if lead and target != lead.status and not converting:
        previous = stages.get(lead.status)
        if previous and target not in previous["allowedNext"]:
            raise HTTPException(
                422, f"Transition from {lead.status} to {target} is not allowed"
            )
    missing = []
    qualification = details.get("qualification") or {}
    for field in stages[target]["required"]:
        value = (
            qualification.get(field)
            if field in {"need", "budget", "decisionMaker", "timeline"}
            else details.get(
                field, values.get(field, getattr(lead, field, None) if lead else None)
            )
        )
        if value is None or value == "" or value is False:
            missing.append(field)
    if target == "Lost" and not details.get("lossReason"):
        missing.append("lossReason")
    if missing:
        raise HTTPException(
            422,
            {
                "message": "Complete the required fields before changing stage",
                "missingFields": sorted(set(missing)),
            },
        )


def assign(db, user, lead, settings, details=None):
    assignment = settings["assignment"]
    owner_ids = assignment.get("ownerIds", [])
    for rule in assignment.get("rules", []):
        matched = (
            any(
                str(item.get("name", "")).casefold() == rule["equals"].casefold()
                for item in (details or {}).get("interests", [])
            )
            if rule["field"] == "product"
            else str(getattr(lead, rule["field"], "") or "").casefold()
            == rule["equals"].casefold()
        )
        if matched:
            check_owner(db, user, rule["ownerId"])
            lead.owner_user_id = rule["ownerId"]
            return
    if owner_ids and assignment["mode"] != "manual":
        # Lock the tenant configuration to serialize round-robin allocation.
        row = (
            db.query(CRMLeadManagementConfig)
            .filter_by(organization_id=lead.organization_id)
            .with_for_update()
            .first()
        )
        for ident in owner_ids:
            check_owner(db, user, ident)
        if assignment["mode"] == "workload":
            counts = dict(
                db.query(CRMLead.owner_user_id, func.count(CRMLead.id))
                .outerjoin(
                    CRMLeadOperatingState,
                    and_(
                        CRMLeadOperatingState.lead_id == CRMLead.id,
                        CRMLeadOperatingState.organization_id == lead.organization_id,
                    ),
                )
                .filter(
                    CRMLead.organization_id == lead.organization_id,
                    CRMLead.id != lead.id,
                    CRMLead.owner_user_id.in_(owner_ids),
                    CRMLead.deleted_at.is_(None),
                    CRMLead.status.notin_(["Converted", "Lost"]),
                    func.coalesce(
                        CRMLeadOperatingState.data_json["archived"].as_boolean(), False
                    ).is_(False),
                )
                .group_by(CRMLead.owner_user_id)
                .all()
            )
            lead.owner_user_id = min(
                owner_ids, key=lambda ident: (counts.get(ident, 0), ident)
            )
        else:
            cursor = int((row.settings_json if row else {}).get("assignmentCursor", 0))
            lead.owner_user_id = owner_ids[cursor % len(owner_ids)]
            if row:
                row.settings_json = {
                    **row.settings_json,
                    "assignmentCursor": cursor + 1,
                }


def workflow(db, lead, user, trigger, settings):
    for rule in settings["workflows"]:
        if not rule.get("enabled", True) or rule["trigger"] != trigger:
            continue
        if rule.get("stage") and lead.status != rule["stage"]:
            continue
        if rule["action"] == "task":
            task(
                db,
                lead,
                user,
                rule["title"],
                now() + timedelta(hours=rule.get("delayHours", 0)),
                "task",
                {"workflow": rule["name"]},
            )
        elif rule["action"] == "notify":
            db.add(
                Notification(
                    company_id=lead.organization_id,
                    user_id=lead.owner_user_id or user.id,
                    module="crm",
                    title=rule["title"],
                    message=f"{rule['name']}: {lead.full_name}",
                    event_type="lead_workflow",
                    action_url=f"/crm/leads?lead={lead.id}",
                )
            )
        elif rule["action"] == "assign":
            check_owner(db, user, rule["ownerId"])
            lead.owner_user_id = rule["ownerId"]
        event(
            db,
            lead,
            user,
            "automation",
            rule["name"],
            {"action": rule["action"], "trigger": trigger},
        )


def task(db, lead, user, title, due, kind="task", metadata=None):
    item = CRMActivity(
        organization_id=lead.organization_id,
        owner_user_id=lead.owner_user_id or user.id,
        entity_type="lead",
        entity_id=lead.id,
        lead_id=lead.id,
        activity_type=kind,
        subject=title[:180],
        title=title[:220],
        status="Pending",
        due_date=due,
        activity_date=due,
        metadata_json=metadata or {},
        created_by_user_id=user.id,
    )
    db.add(item)
    return item


def stop_sequence(db, lead, data, reason):
    sequence = data.get("sequence")
    if sequence and sequence.get("active"):
        sequence = {
            **sequence,
            "active": False,
            "stoppedAt": now().isoformat(),
            "stopReason": reason,
        }
        data["sequence"] = sequence
        rows = (
            db.query(CRMActivity)
            .filter(
                CRMActivity.organization_id == lead.organization_id,
                CRMActivity.lead_id == lead.id,
                CRMActivity.status == "Pending",
            )
            .all()
        )
        for item in rows:
            if (item.metadata_json or {}).get("sequenceId") == sequence["id"]:
                item.status = "Cancelled"


def derived(lead, details, settings, at=None):
    at = at or now()
    created = dt(lead.created_at) or at
    contacted = dt(lead.last_contacted_at)
    follow = dt(lead.next_follow_up_at)
    active = (
        not lead.is_converted and lead.status != "Lost" and not details.get("archived")
    )
    deadline = created + timedelta(hours=settings["slaHours"])
    overdue = bool(active and follow and follow < at)
    sla = bool(active and not contacted and at > deadline)
    last_activity = dt(details.get("lastActivityAt")) or contacted or created
    stale = bool(
        active and (at - last_activity).total_seconds() >= settings["staleDays"] * 86400
    )
    action = (
        "Review qualification and convert"
        if lead.status == "Qualified"
        else "Schedule the next follow-up"
    )
    if not contacted:
        action = "Make first contact"
    if stale:
        action = "Reconnect with this inactive lead"
    if overdue:
        action = "Complete the overdue follow-up"
    if sla:
        action = "Contact now: first-response deadline missed"
    if lead.is_converted:
        action = "Open linked contact and deal"
    elif details.get("archived"):
        action = "Review archive and reactivation date"
    elif lead.status == "Lost":
        action = "Review loss reason and nurture timing"
    return {
        "overdue": overdue,
        "slaBreached": sla,
        "stale": stale,
        "active": active,
        "slaDeadline": deadline.isoformat(),
        "ageDays": max(0, (at - created).days),
        "nextAction": action,
        "priority": (100 if sla else 0)
        + (50 if overdue else 0)
        + (20 if stale else 0)
        + (10 if not contacted and active else 0),
        "followUpLabel": (
            "Overdue"
            if overdue
            else (
                "Due today"
                if follow and follow.date() == at.date()
                else "Upcoming" if follow else "Not scheduled"
            )
        ),
    }


def serialize(db, lead, settings, state_row=None, lookup_state=True):
    row = (
        state_row
        if state_row is not None or not lookup_state
        else state(db, lead, False)
    )
    details = copy.deepcopy(row.data_json if row else {})
    details.pop("booking", None)
    return {
        **core()._serialize(lead),
        "details": details,
        "version": row.version if row else 0,
        "signals": derived(lead, details, settings),
    }


def guard_existing_api(db, user, lead, data):
    """A configured process cannot be bypassed through the classic CRUD endpoints."""
    org = core()._organization_id(db, user)
    if data.get("owner_user_id"):
        check_owner(db, user, data["owner_user_id"])
    if not db.get(CRMLeadManagementConfig, org) and (
        not lead or not state(db, lead, False)
    ):
        return
    settings = config(db, org)
    row = state(db, lead, False) if lead else None
    details = copy.deepcopy(row.data_json if row else {})
    validate_transition(settings, lead, data, details)
    if settings["preventDuplicates"] and (
        lead is None or {"email", "phone"} & data.keys()
    ):
        values = {
            "email": data.get("email", lead.email if lead else None),
            "phone": data.get("phone", lead.phone if lead else None),
        }
        if duplicates(db, org, values, lead.id if lead else None):
            raise HTTPException(
                409, "A matching lead exists. Review duplicates in the lead workbench"
            )
    if "lead_score" in data and not details.get("scoreOverrideReason"):
        raise HTTPException(
            422, "Use the lead workbench to record a reason for the manual score"
        )


def after_existing_api(db, user, lead, before=None, owner_was_provided=False):
    from app.apps.crm.lead_assignment import assignment_changed

    previous_owner = (
        before.get("owner_user_id", lead.owner_user_id) if before is not None else None
    )
    if not db.get(CRMLeadManagementConfig, lead.organization_id) and not state(
        db, lead, False
    ):
        assignment_changed(db, lead, user, previous_owner)
        return
    settings = config(db, lead.organization_id)
    row = state(db, lead)
    data = copy.deepcopy(row.data_json)
    if before is None:
        data["stageEnteredAt"] = now().isoformat()
        data["enquiryInbox"] = lead.source in {
            "Website",
            "Email",
            "WhatsApp",
            "Webform",
        }
        event(db, lead, user, "created", "Lead captured", {"status": lead.status})
        if not owner_was_provided and core()._crm_can_assign_beyond_self(user):
            assign(db, user, lead, settings)
        workflow(db, lead, user, "created", settings)
    elif before.get("status") != lead.status:
        entered = dt(data.get("stageEnteredAt")) or dt(lead.created_at) or now()
        event(
            db,
            lead,
            user,
            "stage_changed",
            f"{before.get('status')} → {lead.status}",
            {
                "from": before.get("status"),
                "to": lead.status,
                "durationHours": round((now() - entered).total_seconds() / 3600, 2),
            },
        )
        data["stageEnteredAt"] = now().isoformat()
        if lead.status in {"Lost", "Converted"}:
            stop_sequence(db, lead, data, lead.status)
        workflow(db, lead, user, "stage_changed", settings)
    row.data_json = data
    row.version += 1
    assignment_changed(db, lead, user, previous_owner)


class Strict(BaseModel):
    model_config = ConfigDict(extra="forbid")


class Stage(Strict):
    name: str = Field(min_length=1, max_length=40)
    required: list[str] = Field(default_factory=list, max_length=20)
    allowedNext: list[str] = Field(default_factory=list, max_length=25)


class AssignmentRule(Strict):
    field: Literal[
        "source", "city", "state", "country", "industry", "territory_id", "product"
    ]
    equals: str = Field(min_length=1, max_length=180)
    ownerId: int = Field(gt=0)


class Assignment(Strict):
    mode: Literal["manual", "round_robin", "workload"] = "manual"
    ownerIds: list[int] = Field(default_factory=list, max_length=100)
    rules: list[AssignmentRule] = Field(default_factory=list, max_length=100)


class SequenceStep(Strict):
    day: int = Field(ge=0, le=365)
    type: Literal["call", "email", "task", "meeting"]
    title: str = Field(min_length=1, max_length=180)


class Sequence(Strict):
    name: str = Field(min_length=1, max_length=120)
    steps: list[SequenceStep] = Field(min_length=1, max_length=30)


class Workflow(Strict):
    name: str = Field(min_length=1, max_length=120)
    enabled: bool = True
    trigger: Literal["created", "stage_changed", "responded"]
    stage: str = ""
    action: Literal["task", "notify", "assign"]
    title: str = Field(min_length=1, max_length=180)
    delayHours: int = Field(default=0, ge=0, le=8760)
    ownerId: int | None = None


class SettingsPayload(Strict):
    stages: list[Stage] = Field(min_length=5, max_length=25)
    slaHours: int = Field(ge=1, le=8760)
    staleDays: int = Field(ge=1, le=365)
    preventDuplicates: bool = True
    assignment: Assignment
    sequences: list[Sequence] = Field(default_factory=list, max_length=30)
    workflows: list[Workflow] = Field(default_factory=list, max_length=100)


class Qualification(Strict):
    need: str = Field(default="", max_length=4000)
    budget: str = Field(default="", max_length=200)
    decisionMaker: str = Field(default="", max_length=200)
    timeline: str = Field(default="", max_length=200)


class Interest(Strict):
    name: str = Field(min_length=1, max_length=180)
    value: Decimal = Field(default=0, ge=0, le=9999999999)
    requirements: rich_text_type(4000, 0) = Field(default="")


class DetailsPatch(Strict):
    qualification: Qualification | None = None
    interests: list[Interest] | None = Field(default=None, max_length=50)
    lossReason: str | None = Field(default=None, max_length=1000)
    competitor: str | None = Field(default=None, max_length=200)
    scoreOverrideReason: str | None = Field(default=None, max_length=1000)
    archived: bool | None = None
    archiveReason: str | None = Field(default=None, max_length=1000)
    reviewAt: datetime | None = None
    assignmentReason: str | None = Field(default=None,max_length=1000)
    customerTimezone: str | None = Field(default=None,max_length=60)

    @field_validator("customerTimezone")
    @classmethod
    def valid_customer_zone(cls,value):
        if not value:return None
        from zoneinfo import ZoneInfo,ZoneInfoNotFoundError
        try:ZoneInfo(value)
        except (ValueError,ZoneInfoNotFoundError):raise ValueError("Choose a valid IANA customer time zone")
        return value


class LeadPayload(Strict):
    values: dict[str, Any] = Field(default_factory=dict)
    details: DetailsPatch = Field(default_factory=DetailsPatch)
    version: int | None = None


def checked_values(values, partial=False):
    # System/audit/conversion attributes are never writable through this surface.
    allowed = set(BASE_TRANSFER_FIELDS) | {
        "first_name",
        "last_name",
        "owner_user_id",
        "lead_score",
        "rating",
        "territory_id",
        "website",
        "tags_text",
        "expected_close_date",
    }
    normalized = core()._normalize_payload(values)
    unknown = set(normalized) - allowed
    if unknown:
        raise HTTPException(
            422, "Unsupported lead fields: " + ", ".join(sorted(unknown))
        )
    if normalized.get("full_name") and not normalized.get("first_name"):
        parts = str(normalized["full_name"]).strip().split(maxsplit=1)
        normalized["first_name"] = parts[0] if parts else ""
        normalized.setdefault("last_name", parts[1] if len(parts) > 1 else None)
    from app.core.input_validation import (
        national_phone,
        calling_code,
        email_address,
        clean_text,
    )

    try:
        if normalized.get("email"):
            normalized["email"] = email_address(normalized["email"])
        if normalized.get("phone"):
            normalized["phone"] = national_phone(normalized["phone"])
            normalized["phone_country_code"] = calling_code(
                normalized.get("phone_country_code")
            )
        for field in ("full_name", "first_name", "last_name", "company_name"):
            if normalized.get(field):
                normalized[field] = clean_text(
                    normalized[field], CRMLead.__table__.columns[field].type.length
                )
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    data = core()._validate_and_build(
        core()._get_resource("leads"), normalized, partial=partial
    )
    for key in ("first_name", "full_name", "status"):
        if key in data and not str(data[key] or "").strip():
            raise HTTPException(422, f"{key} cannot be empty")
    if "lead_score" in data and (
        data["lead_score"] is None or not 0 <= data["lead_score"] <= 100
    ):
        raise HTTPException(422, "Lead score must be between 0 and 100")
    if not partial and not str(data.get("full_name") or "").strip():
        raise HTTPException(422, "Lead name is required")
    for field, value in data.items():
        column = CRMLead.__table__.columns.get(field)
        maximum = getattr(column.type, "length", None) if column is not None else None
        if maximum and isinstance(value, str) and len(value) > maximum:
            raise HTTPException(422, f"{field} exceeds {maximum} characters")
    if "estimated_value" in data and (
        data["estimated_value"] is not None
        and (
            not data["estimated_value"].is_finite()
            or data["estimated_value"] < 0
            or data["estimated_value"] > Decimal("9999999999.99")
        )
    ):
        raise HTTPException(
            422, "Estimated value must be between 0 and 9,999,999,999.99"
        )
    if data.get("notes") and len(data["notes"]) > 32000:
        raise HTTPException(422, "Notes must be 32,000 characters or fewer")
    return data


def create_lead(db, user, payload, settings):
    crm = core()
    crm._assert_entity_permission("leads", user, "manage")
    org = crm._organization_id(db, user)
    data = checked_values(payload.values)
    details = payload.details.model_dump(exclude_unset=True, mode="json")
    if details.get("archived") and not details.get("archiveReason"):
        raise HTTPException(422, "Enter an archive reason")
    validate_transition(settings, None, data, details)
    if settings["preventDuplicates"] and duplicates(db, org, data):
        raise HTTPException(
            409,
            "A lead with this email or phone already exists; review duplicates before creating",
        )
    crm._validate_related_records(db, data, org)
    data.update(
        organization_id=org, created_by_user_id=user.id, updated_by_user_id=user.id
    )
    explicit_owner = bool(data.get("owner_user_id"))
    if not explicit_owner:
        data["owner_user_id"] = user.id
    crm._assert_write_scope(crm._get_resource("leads"), data, user, db)
    lead = CRMLead(**data)
    db.add(lead)
    db.flush()
    if not explicit_owner:
        crm._apply_territory_assignment(
            db, org, lead, override_manual=False, override_owner=True
        )
        # Assignment rules are administered centrally; low-privilege users do not bypass their own write scope.
        if crm._crm_can_assign_beyond_self(user):
            assign(db, user, lead, settings, details)
    if lead.owner_user_id:
        check_owner(db, user, lead.owner_user_id)
    if "lead_score" in data:
        if not details.get("scoreOverrideReason"):
            raise HTTPException(422, "Explain the manual score override")
        lead.lead_score_mode = "manual"
        lead.lead_score_label = crm._score_label(lead.lead_score)
        lead.rating = lead.lead_score_label
    else:
        crm._apply_lead_score(db, lead, org, user.id)
    row = state(db, lead)
    row.data_json = {
        **details,
        "stageEnteredAt": now().isoformat(),
        "enquiryInbox": lead.source in {"Website", "Email", "WhatsApp", "Webform"},
        "enquiryPending": lead.source in {"Website", "Email", "WhatsApp", "Webform"},
    }
    event(db, lead, user, "created", "Lead captured", {"status": lead.status or "New"})
    workflow(db, lead, user, "created", settings)
    from app.apps.crm.lead_assignment import assignment_changed

    assignment_changed(db, lead, user, None)
    return lead


@router.get("/settings")
def get_settings(
    db: Session = Depends(get_db), user: User = Depends(RequirePermission(*ACCESS))
):
    core()._assert_entity_permission("leads", user, "view")
    org = core()._organization_id(db, user)
    settings = config(db, org)
    settings.pop("assignmentCursor", None)
    settings.pop("capture", None)
    settings.pop("productivityViews",None)
    settings.pop("dealStageGuidance",None)
    can_manage = bool(
        user.is_superuser or core()._crm_permission_names(user).intersection(WRITE)
    )
    can_admin = bool(
        user.is_superuser
        or core()._crm_permission_names(user).intersection({"crm_admin"})
    )
    from app.apps.crm.lead_assignment import directory

    owners, unavailable = directory(db, user, core()._crm_can_assign_beyond_self(user))
    return {
        "settings": settings,
        "owners": owners,
        "unavailableEmployees": unavailable,
        "canManage": can_manage,
        "canAdmin": can_admin,
        "canAssign": bool(core()._crm_can_assign_beyond_self(user)),
        "userId": user.id,
    }


@router.put("/settings")
def save_settings(
    payload: SettingsPayload,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission("crm_admin")),
):
    org = core()._organization_id(db, user)
    values = payload.model_dump(mode="json")
    names = [s["name"] for s in values["stages"]]
    if len(names) != len(set(names)) or not {
        "New",
        "Contacted",
        "Qualified",
        "Converted",
        "Lost",
    }.issubset(names):
        raise HTTPException(
            422, "Keep the five system stages and use unique stage names"
        )
    required_allowed = set(TRANSFER_FIELDS) | {
        "need",
        "budget",
        "decisionMaker",
        "timeline",
        "lossReason",
        "competitor",
    }
    for stage in values["stages"]:
        if (
            set(stage["allowedNext"]) - set(names)
            or set(stage["required"]) - required_allowed
        ):
            raise HTTPException(
                422, "Stages contain an unknown transition or required field"
            )
    if len({s["name"] for s in values["sequences"]}) != len(values["sequences"]):
        raise HTTPException(422, "Sequence names must be unique")
    for ident in values["assignment"]["ownerIds"] + [
        r["ownerId"] for r in values["assignment"]["rules"]
    ]:
        check_owner(db, user, ident)
    for rule in values["workflows"]:
        if rule["stage"] and rule["stage"] not in names:
            raise HTTPException(422, "Workflow stage is not configured")
        if rule["action"] == "assign":
            if not rule["ownerId"]:
                raise HTTPException(422, "Assignment workflow needs an owner")
            check_owner(db, user, rule["ownerId"])
    row = db.get(CRMLeadManagementConfig, org)
    if not row:
        row = CRMLeadManagementConfig(organization_id=org)
        db.add(row)
    row.settings_json = {
        **{
            key: value
            for key, value in (row.settings_json or {}).items()
            if key in {"capture", "assignmentCursor", "productivityViews", "dealStageGuidance"}
        },
        **values,
    }
    row.updated_by = user.id
    db.commit()
    return {"settings": values}


class CaptureSettings(Strict):
    enabled: bool
    ownerId: int = Field(gt=0)


@router.get("/capture-settings")
def capture_settings(
    db: Session = Depends(get_db), user: User = Depends(RequirePermission("crm_admin"))
):
    org = core()._organization_id(db, user)
    capture = config(db, org).get("capture", {})
    return {
        "enabled": bool(capture.get("enabled")),
        "path": (
            f"/api/v1/crm/lead-management/capture/{org}/{capture['token']}"
            if capture.get("enabled")
            else None
        ),
        "ownerId": capture.get("ownerId"),
    }


@router.put("/capture-settings")
def enable_capture(
    payload: CaptureSettings,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission("crm_admin")),
):
    org = core()._organization_id(db, user)
    check_owner(db, user, payload.ownerId)
    row = db.get(CRMLeadManagementConfig, org)
    if not row:
        row = CRMLeadManagementConfig(organization_id=org, settings_json={})
        db.add(row)
    # Rotating the link also revokes previous public forms.
    row.settings_json = {
        **(row.settings_json or {}),
        "capture": {
            "enabled": payload.enabled,
            "token": secrets.token_urlsafe(32),
            "ownerId": payload.ownerId,
            "adminId": user.id,
        },
    }
    row.updated_by = user.id
    db.commit()
    return capture_settings(db, user)


def capture_context(db, org, token):
    values = config(db, org)
    capture = values.get("capture", {})
    if not capture.get("enabled") or not secrets.compare_digest(
        str(capture.get("token", "")), token
    ):
        raise HTTPException(404, "This enquiry form is unavailable")
    actor = db.get(User, capture.get("adminId"))
    if not actor or not actor.is_active or core()._organization_id(db, actor) != org:
        raise HTTPException(404, "This enquiry form is unavailable")
    if not actor.is_superuser and "crm_admin" not in core()._crm_permission_names(
        actor
    ):
        raise HTTPException(404, "This enquiry form is unavailable")
    return values, capture, actor


@router.get("/capture/{org}/{token}", response_class=HTMLResponse)
def public_capture_form(org: int, token: str, db: Session = Depends(get_db)):
    capture_context(db, org, token)
    return public_page(
        "Send an enquiry",
        '<p>Tell us what you need. Our team will follow up with you.</p><form method="post"><label>Your name<input name="full_name" maxlength="220" required autocomplete="name"></label><label>Email<input type="email" name="email" maxlength="150" required autocomplete="email"></label><label>Company<input name="company_name" maxlength="180" autocomplete="organization"></label><label>How can we help?<textarea name="message" maxlength="4000" rows="4" required></textarea></label><label class="trap" aria-hidden="true">Leave empty<input name="website" tabindex="-1" autocomplete="off"></label><label><input name="consent" type="checkbox" value="yes" required> I agree to be contacted about this enquiry.</label><button>Submit enquiry</button></form>',
    )


@router.post("/capture/{org}/{token}")
def submit_capture(
    org: int,
    token: str,
    full_name: str = Form(max_length=220),
    email: str = Form(max_length=150),
    company_name: str = Form(default="", max_length=180),
    message: str = Form(max_length=4000),
    website: str = Form(default="", max_length=500),
    consent: str = Form(default=""),
    db: Session = Depends(get_db),
):
    settings, capture, actor = capture_context(db, org, token)
    if website:
        return public_page("Thank you", "<p>Your enquiry has been received.</p>")
    if consent != "yes":
        raise HTTPException(422, "Please agree to be contacted about this enquiry")
    # A tenant-owned database lock makes the hourly intake cap consistent across workers.
    db.query(CRMLeadManagementConfig).filter_by(
        organization_id=org
    ).with_for_update().first()
    received = (
        db.query(CRMTimelineEvent)
        .filter(
            CRMTimelineEvent.organization_id == org,
            CRMTimelineEvent.event_type == "web_enquiry",
            CRMTimelineEvent.occurred_at >= now() - timedelta(hours=1),
        )
        .count()
    )
    if received >= 100:
        raise HTTPException(
            429, "This form has received too many enquiries. Please try again later"
        )
    values = {
        "full_name": full_name,
        "email": email,
        "company_name": company_name or None,
        "source": "Website",
        "notes": message,
    }
    checked_values(values)
    existing_ids = duplicates(db, org, values)
    if not existing_ids:
        # Keep the responsible intake owner; assignment automation can be configured separately.
        values["owner_user_id"] = capture["ownerId"]
        lead = create_lead(
            db,
            actor,
            LeadPayload(
                values=values,
                details=DetailsPatch(qualification=Qualification(need=message)),
            ),
            settings,
        )
        row = state(db, lead)
        row.data_json = {
            **row.data_json,
            "enquiryInbox": True,
            "enquiryPending": True,
            "contactConsentAt": now().isoformat(),
        }
        event(
            db,
            lead,
            actor,
            "web_enquiry",
            "Website enquiry received",
            {"consent": True},
        )
    else:
        lead = (
            db.query(CRMLead)
            .filter(CRMLead.id == existing_ids[0], CRMLead.organization_id == org)
            .with_for_update()
            .one()
        )
        row = state(db, lead)
        row.data_json = {
            **row.data_json,
            "enquiryInbox": True,
            "enquiryPending": True,
            "contactConsentAt": now().isoformat(),
        }
        row.version += 1
        event(
            db,
            lead,
            actor,
            "web_enquiry",
            "Another website enquiry received",
            {"consent": True},
            body=message,
        )
    db.commit()
    return public_page(
        "Thank you", "<p>Your enquiry has been received. Our team will contact you.</p>"
    )


@router.get("/leads")
def list_leads(
    search: str = "",
    stage: str = "",
    owner: int | None = None,
    queue: Literal[
        "all",
        "my_day",
        "active",
        "sla",
        "untouched",
        "overdue",
        "stale",
        "archived",
        "unassigned",
        "reactivation",
    ] = "all",
    page: int = Query(1, ge=1),
    per_page: int = Query(40, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*ACCESS)),
):
    settings = config(db, core()._organization_id(db, user))
    query = visible_query(db, user)
    if search:
        query = query.filter(
            or_(
                *[
                    getattr(CRMLead, key).ilike(f"%{search}%")
                    for key in ["full_name", "email", "phone", "company_name"]
                ]
            )
        )
    if stage:
        query = query.filter(CRMLead.status == stage)
    if owner is not None:
        query = query.filter(CRMLead.owner_user_id == owner)
    leads = query.order_by(CRMLead.id.desc()).limit(MAX_ROWS + 1).all()
    truncated = len(leads) > MAX_ROWS
    leads = leads[:MAX_ROWS]
    states = (
        {
            s.lead_id: s
            for s in db.query(CRMLeadOperatingState)
            .filter(
                CRMLeadOperatingState.organization_id
                == core()._organization_id(db, user),
                CRMLeadOperatingState.lead_id.in_([l.id for l in leads]),
            )
            .all()
        }
        if leads
        else {}
    )
    items = [
        serialize(db, lead, settings, states.get(lead.id), lookup_state=False)
        for lead in leads
    ]
    at = now()

    def match(item):
        details, signal = item["details"], item["signals"]
        if queue == "archived":
            return bool(details.get("archived"))
        if queue == "reactivation":
            return bool(
                details.get("reviewAt")
                and dt(details["reviewAt"]) <= at
                and not item.get("is_converted")
            )
        if details.get("archived"):
            return False
        if queue == "active":
            return signal["active"]
        if queue == "sla":
            return signal["slaBreached"]
        if queue == "my_day":
            return (
                item.get("owner_user_id") == user.id
                and signal["active"]
                and (
                    signal["priority"] > 0
                    or bool(
                        item.get("next_follow_up_at")
                        and dt(item["next_follow_up_at"]).date() <= at.date()
                    )
                )
            )
        if queue == "untouched":
            return signal["active"] and not item.get("last_contacted_at")
        if queue in {"overdue", "stale"}:
            return signal[queue]
        if queue == "unassigned":
            return not item.get("owner_user_id") or bool(
                details.get(
                    "enquiryPending",
                    details.get("enquiryInbox") and not item.get("last_contacted_at"),
                )
            )
        return True

    filtered = [item for item in items if match(item)]
    if queue == "my_day":
        filtered.sort(
            key=lambda item: (
                -item["signals"]["priority"],
                item.get("next_follow_up_at") or item.get("created_at") or "",
            )
        )
    return {
        "items": filtered[(page - 1) * per_page : page * per_page],
        "total": len(filtered),
        "page": page,
        "perPage": per_page,
        "truncated": truncated,
        "limit": MAX_ROWS,
        "summary": {
            "total": len(items),
            "active": sum(i["signals"]["active"] for i in items),
            "overdue": sum(i["signals"]["overdue"] for i in items),
            "slaBreached": sum(i["signals"]["slaBreached"] for i in items),
            "stale": sum(i["signals"]["stale"] for i in items),
        },
        "byStage": dict(Counter(i["status"] for i in items)),
        "bySource": dict(Counter(i.get("source") or "Other" for i in items)),
        "convertedBySource": dict(
            Counter(i.get("source") or "Other" for i in items if i.get("is_converted"))
        ),
        "lostReasons": dict(
            Counter(
                i["details"].get("lossReason") or "Not recorded"
                for i in items
                if i["status"] == "Lost"
            )
        ),
    }


@router.post("/leads", status_code=201)
def add_lead(
    payload: LeadPayload,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    settings = config(db, core()._organization_id(db, user))
    lead = create_lead(db, user, payload, settings)
    db.commit()
    return serialize(db, lead, settings)


@router.get("/leads/{ident}")
def detail(
    ident: int,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*ACCESS)),
):
    lead = lead_record(db, user, ident)
    result = serialize(db, lead, config(db, lead.organization_id))
    events = (
        db.query(CRMTimelineEvent)
        .filter_by(
            organization_id=lead.organization_id, record_type="lead", record_id=lead.id
        )
        .order_by(CRMTimelineEvent.occurred_at.desc(), CRMTimelineEvent.id.desc())
        .limit(100)
        .all()
    )
    result["timeline"] = [core()._serialize(e) for e in events]
    from .productivity import communication_history
    result["timeline"] += [{"id":-e.id,"event_type":f"{e.channel}_{e.event_type}","title":e.subject or f"{e.channel}: {e.event_type}","description":e.summary,"occurred_at":e.created_at,"actor_user_id":e.actor_user_id,"metadata_json":e.metadata_json} for e in communication_history(db,user,"leads",lead.id,lead.organization_id)]
    result["timeline"].sort(key=lambda e:str(e.get("occurred_at") or e.get("occurredAt") or ""),reverse=True)
    crm=core()
    manages=bool(user.is_superuser or crm._crm_permission_names(user).intersection(WRITE))
    grant=(getattr(user,"_suite_grants",None) or {}).get("crm")
    exports=grant is None or "export" in grant.get("actions",[])
    result["permissions"]={"edit":manages,"assign":crm._crm_can_assign_beyond_self(user),"export":exports,"convert":manages and lead.status=="Qualified" and not lead.is_converted,"edit_reason":"" if manages else "CRM lead management permission is required","assign_reason":"" if crm._crm_can_assign_beyond_self(user) else "Only permitted sales managers and administrators can assign another salesperson","export_reason":"" if exports else "Your organization access policy does not permit export","convert_reason":"" if manages and lead.status=="Qualified" and not lead.is_converted else "Qualify this lead and obtain CRM lead management access before converting; already converted leads cannot be converted again"}
    result["activities"] = [
        core()._serialize(a)
        for a in db.query(CRMActivity)
        .filter_by(organization_id=lead.organization_id, lead_id=lead.id)
        .order_by(CRMActivity.id.desc())
        .limit(100)
        .all()
    ]
    _, _, reasons = core()._calculate_lead_score(
        db, lead, lead.organization_id, user.id
    )
    result["scoreReasons"] = reasons
    # Default scoring rules may be provisioned on first use.
    db.commit()
    return result


def apply_patch(db, user, lead, payload, settings):
    row = state(db, lead, False)
    if payload.version is not None and payload.version != (row.version if row else 0):
        raise HTTPException(
            409, "This lead changed since you opened it. Refresh before saving"
        )
    row = row or state(db, lead)
    data = checked_values(payload.values, True)
    details = {
        **copy.deepcopy(row.data_json),
        **payload.details.model_dump(exclude_unset=True, mode="json"),
    }
    validate_transition(settings, lead, data, details)
    if (
        settings["preventDuplicates"]
        and ({"email", "phone", "phone_country_code"} & data.keys())
        and duplicates(
            db,
            lead.organization_id,
            {
                "email": data.get("email", lead.email),
                "phone": data.get("phone", lead.phone),
                "phone_country_code": data.get(
                    "phone_country_code", lead.phone_country_code
                ),
            },
            lead.id,
        )
    ):
        raise HTTPException(409, "Email or phone matches another lead")
    if "lead_score" in data and not details.get("scoreOverrideReason"):
        raise HTTPException(422, "Explain the manual score override")
    if "lead_score" in data and not 0 <= data["lead_score"] <= 100:
        raise HTTPException(422, "Lead score must be between 0 and 100")
    if data.get("owner_user_id"):
        check_owner(db, user, data["owner_user_id"])
    core()._assert_write_scope(
        core()._get_resource("leads"),
        {"owner_user_id": data.get("owner_user_id", lead.owner_user_id)},
        user,
        db,
    )
    core()._validate_related_records(db, data, lead.organization_id)
    before = {key: core()._json_ready(getattr(lead, key)) for key in data}
    previous_owner = lead.owner_user_id
    old_stage = lead.status
    for key, value in data.items():
        setattr(lead, key, value)
    if "lead_score" in data:
        lead.lead_score_mode = "manual"
        lead.lead_score_label = core()._score_label(lead.lead_score)
        lead.rating = lead.lead_score_label
    else:
        core()._apply_lead_score(db, lead, lead.organization_id, user.id)
    if old_stage != lead.status:
        entered = dt(details.get("stageEnteredAt")) or dt(lead.created_at) or now()
        event(
            db,
            lead,
            user,
            "stage_changed",
            f"{old_stage} → {lead.status}",
            {
                "from": old_stage,
                "to": lead.status,
                "durationHours": round((now() - entered).total_seconds() / 3600, 2),
            },
        )
        details["stageEnteredAt"] = now().isoformat()
        if lead.status in {"Lost", "Converted"}:
            stop_sequence(db, lead, details, lead.status)
        workflow(db, lead, user, "stage_changed", settings)
    if details.get("archived"):
        if not details.get("archiveReason"):
            raise HTTPException(422, "Enter an archive reason")
        stop_sequence(db, lead, details, "Archived")
    event(
        db,
        lead,
        user,
        "updated",
        "Lead updated",
        {
            "before": before,
            "fields": sorted(data),
            "details": sorted(payload.details.model_fields_set),
        },
    )
    lead.updated_by_user_id = user.id
    lead.updated_at = now()
    row.data_json = details
    row.version += 1
    from app.apps.crm.lead_assignment import assignment_changed

    assignment_changed(db, lead, user, previous_owner, reason=payload.details.assignmentReason)
    return lead


@router.patch("/leads/{ident}")
def update_lead(
    ident: int,
    payload: LeadPayload,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    lead = lead_record(db, user, ident, True, True)
    settings = config(db, lead.organization_id)
    apply_patch(db, user, lead, payload, settings)
    db.commit()
    return serialize(db, lead, settings)


class FollowUpPayload(Strict):
    title: str = Field(min_length=1, max_length=180)
    dueAt: datetime
    type: Literal["call", "email", "meeting", "task"] = "task"
    ownerId: int | None = Field(None, gt=0)
    activityId: int | None = Field(None, gt=0)
    rescheduleReason: str | None = Field(default=None,max_length=1000)
    version: int = Field(ge=0)


def available_teammates(db, actor, lead):
    from app.apps.crm.lead_assignment import directory

    owners, _ = directory(db, actor, True)
    permitted = []
    for owner in owners:
        candidate = db.get(User, owner["id"])
        try:
            from app.access.service import background_scope
            with background_scope(db, candidate, "crm", "view"):
                if visible_query(db, candidate).filter(CRMLead.id == lead.id).first():
                    permitted.append(owner)
        except HTTPException:
            continue
    return permitted


@router.get("/leads/{ident}/team")
def lead_team(
    ident: int,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*ACCESS)),
):
    lead = lead_record(db, user, ident)
    return {"members": available_teammates(db, user, lead)}


def check_version(row, version):
    if version != (row.version if row else 0):
        raise HTTPException(
            409,
            "This lead changed. Refresh before saving; your draft is still available",
        )


def refresh_next_follow_up(db, lead):
    db.flush()
    next_task = (
        db.query(CRMActivity)
        .filter(
            CRMActivity.organization_id == lead.organization_id,
            CRMActivity.lead_id == lead.id,
            CRMActivity.status == "Pending",
            CRMActivity.due_date.is_not(None),
        )
        .order_by(CRMActivity.due_date, CRMActivity.id)
        .first()
    )
    lead.next_follow_up_at = next_task.due_date if next_task else None


@router.post("/leads/{ident}/follow-up")
def schedule_follow_up(
    ident: int,
    payload: FollowUpPayload,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    lead = lead_record(db, user, ident, True, True)
    check_version(state(db, lead, False), payload.version)
    if (
        lead.is_converted
        or lead.status == "Lost"
        or (state(db, lead, False) and state(db, lead, False).data_json.get("archived"))
    ):
        raise HTTPException(409, "Reopen the lead before scheduling follow-ups")
    if not payload.title.strip():
        raise HTTPException(422, "Enter a follow-up title")
    due = dt(payload.dueAt)
    if due <= now():
        raise HTTPException(422, "Choose a future follow-up time")
    owner_id = payload.ownerId or lead.owner_user_id or user.id
    check_owner(db, user, owner_id)
    if owner_id not in {member["id"] for member in available_teammates(db, user, lead)}:
        raise HTTPException(422, "This teammate does not have access to this lead")
    if payload.activityId:
        item = (
            db.query(CRMActivity)
            .filter_by(
                id=payload.activityId,
                organization_id=lead.organization_id,
                lead_id=lead.id,
            )
            .with_for_update()
            .first()
        )
        if not item or item.status != "Pending":
            raise HTTPException(409, "This follow-up is no longer pending")
        if not payload.rescheduleReason or len(payload.rescheduleReason.strip())<3:raise HTTPException(422,"Explain why this follow-up is being rescheduled")
        previous_due=item.due_date
        item.subject = item.title = payload.title.strip()
        item.due_date = item.activity_date = due
        item.activity_type = payload.type
        item.owner_user_id = owner_id
        metadata = copy.deepcopy(item.metadata_json or {})
        metadata.pop("followUpReminderAt", None)
        metadata["commitmentHistory"]=[*metadata.get("commitmentHistory",[]),{"previous":str(previous_due),"new":due.isoformat(),"reason":payload.rescheduleReason,"actor":user.id}]
        item.metadata_json = metadata
    else:
        item = task(
            db,
            lead,
            user,
            payload.title.strip(),
            due,
            payload.type,
            {"manualFollowUp": True},
        )
        item.owner_user_id = owner_id
    row = state(db, lead)
    row.version += 1
    lead.updated_at = now()
    refresh_next_follow_up(db, lead)
    event(
        db,
        lead,
        user,
        "follow_up_scheduled",
        "Follow-up rescheduled" if payload.activityId else "Follow-up scheduled",
        {"activityId": item.id, "ownerId": owner_id, "dueAt": due.isoformat(),"previousCommitment":str(previous_due) if payload.activityId else None,"reason":payload.rescheduleReason},
        payload.title.strip(),
    )
    if owner_id != user.id:
        db.add(
            Notification(
                company_id=lead.organization_id,
                user_id=owner_id,
                module="crm",
                event_type="lead_follow_up",
                title="Lead follow-up assigned",
                message=f"{lead.full_name}: {payload.title.strip()}",
                related_entity_type="lead",
                related_entity_id=lead.id,
                action_url=f"/crm/leads?lead={lead.id}",
            )
        )
    db.commit()
    return serialize(db, lead, config(db, lead.organization_id))


class TeamNotePayload(Strict):
    body: rich_text_type(4000, 1) = Field()
    mentionIds: list[int] = Field(default_factory=list, max_length=20)
    version: int = Field(ge=0)


@router.post("/leads/{ident}/team-notes")
def add_team_note(
    ident: int,
    payload: TeamNotePayload,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    lead = lead_record(db, user, ident, True, True)
    check_version(state(db, lead, False), payload.version)
    if not payload.body.strip():
        raise HTTPException(422, "Write a team update")
    members = {member["id"]: member for member in available_teammates(db, user, lead)}
    ids = list(dict.fromkeys(payload.mentionIds))
    if any(ident not in members for ident in ids):
        raise HTTPException(422, "Mention only teammates who have access to this lead")
    event(
        db,
        lead,
        user,
        "team_note",
        "Team update",
        {
            "mentionIds": ids,
            "authorName": members.get(user.id, {}).get("name", user.email),
        },
        payload.body.strip(),
    )
    for recipient in ids:
        if recipient != user.id:
            db.add(
                Notification(
                    company_id=lead.organization_id,
                    user_id=recipient,
                    module="crm",
                    event_type="lead_team_mention",
                    title="You were mentioned on a lead",
                    message=f"Team update on {lead.full_name}",
                    related_entity_type="lead",
                    related_entity_id=lead.id,
                    action_url=f"/crm/leads?lead={lead.id}",
                )
            )
    state(db, lead).version += 1
    lead.updated_at = now()
    db.commit()
    return {"saved": True}


class ActivityPayload(Strict):
    type: Literal["call", "email", "meeting", "note", "task"]
    outcome: str = Field(min_length=1, max_length=4000)
    nextFollowUp: datetime | None = None
    responded: bool = False
    startAt: datetime | None = None


@router.post("/leads/{ident}/activity")
def log_activity(
    ident: int,
    payload: ActivityPayload,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    lead = lead_record(db, user, ident, True, True)
    row = state(db, lead)
    data = copy.deepcopy(row.data_json)
    data["lastActivityAt"] = now().isoformat()
    if payload.type == "meeting" and payload.startAt:
        meeting_start = dt(payload.startAt)
        if meeting_start < now():
            raise HTTPException(422, "Choose a future meeting time")
        meeting_end = meeting_start + timedelta(minutes=30)
        meeting_owner = lead.owner_user_id or user.id
        db.query(User).filter(User.id == meeting_owner).with_for_update().first()
        if (
            db.query(CRMMeeting)
            .filter(
                CRMMeeting.organization_id == lead.organization_id,
                CRMMeeting.owner_user_id == meeting_owner,
                CRMMeeting.deleted_at.is_(None),
                CRMMeeting.status != "Cancelled",
                CRMMeeting.start_time < meeting_end,
                CRMMeeting.end_time > meeting_start,
            )
            .first()
        ):
            raise HTTPException(
                409, "The assigned owner already has a meeting at this time"
            )
        db.add(
            CRMMeeting(
                organization_id=lead.organization_id,
                owner_user_id=meeting_owner,
                lead_id=lead.id,
                title=payload.outcome[:180],
                start_time=meeting_start,
                end_time=meeting_end,
                status="Scheduled",
                created_by_user_id=user.id,
            )
        )
        event(
            db,
            lead,
            user,
            "meeting_scheduled",
            "Meeting scheduled",
            {"startAt": payload.startAt.isoformat()},
            payload.outcome,
        )
    else:
        item = core()._create_timeline_activity(
            db,
            organization_id=lead.organization_id,
            entity_type="lead",
            entity_id=lead.id,
            activity_type=payload.type,
            title=f"{payload.type.title()} outcome",
            body=payload.outcome,
            user_id=user.id,
            metadata={"responded": payload.responded},
        )
        item.status = "Completed"
        item.completed_at = now()
        item.outcome = payload.outcome
        if payload.type in {"call", "email", "meeting"}:
            data.setdefault("firstContactAt", now().isoformat())
            data["enquiryPending"] = False
            lead.last_contacted_at = now()
    if payload.nextFollowUp:
        lead.next_follow_up_at = dt(payload.nextFollowUp)
        task(
            db,
            lead,
            user,
            "Follow up: " + payload.outcome[:150],
            dt(payload.nextFollowUp),
        )
    if payload.responded:
        stop_sequence(db, lead, data, "Lead responded")
        workflow(db, lead, user, "responded", config(db, lead.organization_id))
    row.data_json = data
    row.version += 1
    lead.updated_at = now()
    core()._apply_lead_score(db, lead, lead.organization_id, user.id)
    db.commit()
    return {"saved": True}


@router.post("/leads/{ident}/activities/{activity_id}/complete")
def complete_lead_task(
    ident: int,
    activity_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    lead = lead_record(db, user, ident, True, True)
    activity = (
        db.query(CRMActivity)
        .filter_by(
            id=activity_id, organization_id=lead.organization_id, lead_id=lead.id
        )
        .with_for_update()
        .first()
    )
    if not activity:
        raise HTTPException(404, "This activity was not found for the lead")
    if activity.status != "Pending":
        raise HTTPException(409, "This activity is already completed or cancelled")
    activity.status = "Completed"
    activity.completed_at = now()
    row = state(db, lead)
    data = copy.deepcopy(row.data_json)
    data["lastActivityAt"] = now().isoformat()
    if activity.activity_type in {"call", "email", "meeting"}:
        lead.last_contacted_at = now()
        data["enquiryPending"] = False
        data.setdefault("firstContactAt", now().isoformat())
    db.flush()
    pending = (
        db.query(CRMActivity)
        .filter(
            CRMActivity.organization_id == lead.organization_id,
            CRMActivity.lead_id == lead.id,
            CRMActivity.status == "Pending",
        )
        .order_by(CRMActivity.due_date, CRMActivity.id)
        .all()
    )
    lead.next_follow_up_at = pending[0].due_date if pending else None
    sequence = data.get("sequence")
    if (
        sequence
        and sequence.get("active")
        and not any(
            (a.metadata_json or {}).get("sequenceId") == sequence["id"] for a in pending
        )
    ):
        data["sequence"] = {
            **sequence,
            "active": False,
            "stopReason": "All steps completed",
            "stoppedAt": now().isoformat(),
        }
    row.data_json = data
    row.version += 1
    event(
        db, lead, user, "task_completed", activity.subject, {"activityId": activity.id}
    )
    db.commit()
    return {
        "completed": True,
        "nextFollowUp": core()._json_ready(lead.next_follow_up_at),
    }


@router.post("/leads/{ident}/booking-link")
def booking_link(
    ident: int,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    lead = lead_record(db, user, ident, True, True)
    if not derived(lead, state(db, lead).data_json, config(db, lead.organization_id))[
        "active"
    ]:
        raise HTTPException(422, "Reactivate this lead before offering a meeting")
    row = state(db, lead)
    token = secrets.token_urlsafe(32)
    expires = now() + timedelta(days=30)
    row.data_json = {
        **row.data_json,
        "booking": {"token": token, "expires": expires.isoformat(), "used": False},
    }
    row.version += 1
    event(db, lead, user, "booking_link", "Meeting booking link created")
    db.commit()
    return {
        "path": f"/api/v1/crm/lead-management/booking/{lead.organization_id}/{lead.id}/{token}",
        "expires": expires.isoformat(),
    }


def booking_context(db, org, ident, token):
    lead = (
        db.query(CRMLead)
        .filter_by(id=ident, organization_id=org)
        .filter(CRMLead.deleted_at.is_(None))
        .with_for_update()
        .first()
    )
    row = state(db, lead, False) if lead else None
    booking = (row.data_json if row else {}).get("booking", {})
    if (
        not booking
        or not secrets.compare_digest(str(booking.get("token", "")), token)
        or booking.get("used")
        or dt(booking.get("expires")) < now()
    ):
        raise HTTPException(404, "This booking link is unavailable or has expired")
    if not derived(lead, row.data_json, config(db, org))["active"]:
        raise HTTPException(404, "This booking link is unavailable")
    owner = db.get(User, lead.owner_user_id) if lead.owner_user_id else None
    if not owner or not owner.is_active or core()._organization_id(db, owner) != org:
        raise HTTPException(404, "This booking link is unavailable")
    return lead, row, owner


@router.get("/booking/{org}/{ident}/{token}")
def booking_form(org: int, ident: int, token: str, db: Session = Depends(get_db)):
    booking_context(db, org, ident, token)
    return public_page(
        "Book a conversation",
        '<p>Choose a time for a 30-minute meeting. Your booking is confirmed only if the time is available.</p><form method="post"><label>Meeting date and time<input name="start_time" type="datetime-local" required></label><label>Time zone<select name="offset"><option value="+05:30">India (UTC +05:30)</option><option value="+00:00">UTC</option><option value="+04:00">UAE (UTC +04:00)</option><option value="+08:00">Singapore (UTC +08:00)</option></select></label><button>Confirm meeting</button></form>',
    )


@router.post("/booking/{org}/{ident}/{token}")
def book_meeting(
    org: int,
    ident: int,
    token: str,
    start_time: str = Form(max_length=32),
    offset: str = Form(max_length=6),
    db: Session = Depends(get_db),
):
    lead, row, owner = booking_context(db, org, ident, token)
    if offset not in {"+05:30", "+00:00", "+04:00", "+08:00"} or not re.fullmatch(
        r"\d{4}-\d\d-\d\dT\d\d:\d\d", start_time
    ):
        raise HTTPException(422, "Choose a valid meeting time and time zone")
    start = dt(start_time + offset)
    if not now() + timedelta(minutes=15) < start < now() + timedelta(days=90):
        raise HTTPException(
            422, "Choose a time at least 15 minutes from now and within 90 days"
        )
    end = start + timedelta(minutes=30)
    db.query(User).filter(User.id == owner.id).with_for_update().first()
    if (
        db.query(CRMMeeting)
        .filter(
            CRMMeeting.organization_id == org,
            CRMMeeting.owner_user_id == owner.id,
            CRMMeeting.deleted_at.is_(None),
            CRMMeeting.status != "Cancelled",
            CRMMeeting.start_time < end,
            CRMMeeting.end_time > start,
        )
        .first()
    ):
        response = public_page(
            "Choose another time",
            "<p>This time is already booked. Return to the previous page and choose another time.</p>",
        )
        response.status_code = 409
        return response
    db.add(
        CRMMeeting(
            organization_id=org,
            owner_user_id=owner.id,
            lead_id=lead.id,
            title=f"Conversation with {lead.full_name}"[:180],
            start_time=start,
            end_time=end,
            status="Scheduled",
            created_by_user_id=owner.id,
        )
    )
    data = copy.deepcopy(row.data_json)
    data["booking"]["used"] = True
    data["enquiryPending"] = False
    stop_sequence(db, lead, data, "Meeting booked")
    row.data_json = data
    row.version += 1
    lead.next_follow_up_at = start
    lead.last_contacted_at = now()
    event(
        db,
        lead,
        owner,
        "meeting_booked",
        "Meeting booked by lead",
        {"startAt": start.isoformat()},
    )
    db.add(
        Notification(
            company_id=org,
            user_id=owner.id,
            module="crm",
            title="Lead booked a meeting",
            message=f"{lead.full_name} booked {start_time} (UTC {offset})",
            event_type="lead_booking",
            action_url=f"/crm/leads?lead={lead.id}",
        )
    )
    db.commit()
    return public_page(
        "Meeting confirmed",
        "<p>Your 30-minute meeting is booked for "
        + html.escape(start_time.replace("T", " ") + " (UTC " + offset + ")")
        + ".</p>",
    )


class SequencePayload(Strict):
    name: str = Field(max_length=120)
    stop: bool = False


@router.post("/leads/{ident}/sequence")
def enroll_sequence(
    ident: int,
    payload: SequencePayload,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    from uuid import uuid4

    lead = lead_record(db, user, ident, True, True)
    row = state(db, lead)
    data = copy.deepcopy(row.data_json)
    stop_sequence(db, lead, data, "Stopped" if payload.stop else "Replaced")
    if not payload.stop:
        if lead.is_converted or lead.status == "Lost" or data.get("archived"):
            raise HTTPException(422, "Reactivate this lead before starting a sequence")
        sequence = next(
            (
                s
                for s in config(db, lead.organization_id)["sequences"]
                if s["name"] == payload.name
            ),
            None,
        )
        if not sequence:
            raise HTTPException(422, "Choose a configured sequence")
        ident_sequence = str(uuid4())
        data["sequence"] = {
            "id": ident_sequence,
            "name": payload.name,
            "active": True,
            "startedAt": now().isoformat(),
        }
        due_dates = []
        for step in sequence["steps"]:
            due = now() + timedelta(days=step["day"])
            due_dates.append(due)
            task(
                db,
                lead,
                user,
                step["title"],
                due,
                step["type"],
                {"sequenceId": ident_sequence, "sequence": payload.name},
            )
        lead.next_follow_up_at = min(due_dates)
    row.data_json = data
    row.version += 1
    event(
        db,
        lead,
        user,
        "sequence",
        "Sequence stopped" if payload.stop else f"Started {payload.name}",
    )
    db.commit()
    return {"sequence": data.get("sequence")}


class BulkPayload(Strict):
    ids: list[int] = Field(min_length=1, max_length=200)
    patch: LeadPayload


@router.get("/leads/{ident}/conversion-preview")
def conversion_preview(
    ident: int,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    lead = lead_record(db, user, ident, True)
    crm = core()
    matches = {}
    for entity, model, field, value in [
        ("contacts", CRMContact, "email", lead.email),
        ("companies", CRMCompany, "name", lead.company_name),
    ]:
        resource = crm._get_resource(entity)
        query = crm._apply_record_visibility(
            crm._base_query(db, resource, lead.organization_id), resource, user, db
        )
        matches[entity] = (
            [
                {
                    "id": item.id,
                    "name": getattr(item, "full_name", None)
                    or getattr(item, "name", None),
                }
                for item in query.filter(getattr(model, field) == value).limit(20).all()
            ]
            if value
            else []
        )
    return {
        "leadId": lead.id,
        "leadName": lead.full_name,
        "matches": matches,
        "creates": ["Contact", "Company" if lead.company_name else None, "Deal"],
        "alreadyConverted": bool(lead.is_converted),
        "ready": lead.status == "Qualified" and not lead.is_converted,
    }


class ConversionPayload(Strict):
    existingContactId: int | None = None
    existingCompanyId: int | None = None
    dealName: str | None = Field(default=None, max_length=180)
    createDeal: bool = True


@router.post("/leads/{ident}/convert")
def convert(
    ident: int,
    payload: ConversionPayload,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    crm = core()
    lead = lead_record(db, user, ident, True, True)
    if lead.status != "Qualified":
        raise HTTPException(422, "Qualify the lead before converting it")
    row = state(db, lead)
    settings = config(db, lead.organization_id)
    validate_transition(
        settings, lead, {"status": "Converted"}, row.data_json, converting=True
    )
    # Require explicit duplicate reuse instead of silently creating another contact.
    if lead.email and not payload.existingContactId:
        match = (
            db.query(CRMContact.id)
            .filter(
                CRMContact.organization_id == lead.organization_id,
                CRMContact.deleted_at.is_(None),
                func.lower(CRMContact.email) == lead.email.lower(),
            )
            .first()
        )
        if match:
            raise HTTPException(
                409, "A matching contact exists. Select it in the conversion preview"
            )
    result = crm.convert_lead(
        ident, crm.CRMLeadConvertPayload(**payload.model_dump()), db, user
    )
    data = copy.deepcopy(row.data_json)
    stop_sequence(db, lead, data, "Converted")
    data["stageEnteredAt"] = now().isoformat()
    row.data_json = data
    row.version += 1
    db.commit()
    return result


@router.post("/bulk/preview")
def bulk_preview(
    payload: BulkPayload,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    leads = [lead_record(db, user, ident, True) for ident in set(payload.ids)]
    owner_id=payload.patch.values.get("owner_user_id")
    workload=None
    if owner_id:
        check_owner(db,user,int(owner_id))
        workload={"owner_id":int(owner_id),"visible_active_leads":visible_query(db,user).filter(CRMLead.owner_user_id==int(owner_id),CRMLead.is_converted.is_(False),CRMLead.status!="Lost").count(),"selected_leads":len(leads),"pending_followups":db.query(CRMActivity).filter(CRMActivity.organization_id==core()._organization_id(db,user),CRMActivity.lead_id.in_([l.id for l in leads]),CRMActivity.status.in_(["Pending","To Do","Open"])).count()}
    return {"workload":workload,"organization_id":core()._organization_id(db,user),
        "count": len(leads),
        "items": [{"id": l.id, "name": l.full_name, "stage": l.status} for l in leads],
        "patch": payload.patch.model_dump(mode="json"),
    }


@router.post("/bulk/apply")
def bulk_apply(
    payload: BulkPayload,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    settings = config(db, core()._organization_id(db, user))
    with db.begin_nested():
        leads = [
            lead_record(db, user, ident, True, True)
            for ident in sorted(set(payload.ids))
        ]
        for lead in leads:
            apply_patch(db, user, lead, payload.patch, settings)
    db.commit()
    return {"updated": len(leads)}


@router.post("/alerts/evaluate")
def evaluate_alerts(
    db: Session = Depends(get_db), user: User = Depends(RequirePermission(*WRITE))
):
    """Idempotent in-app escalation; may also be called by an authenticated scheduler."""
    settings = config(db, core()._organization_id(db, user))
    count = 0
    for lead in (
        visible_query(db, user)
        .limit(MAX_ROWS)
        .with_for_update(of=CRMLead, skip_locked=True)
        .all()
    ):
        row = state(db, lead)
        data = copy.deepcopy(row.data_json)
        signals = derived(lead, data, settings)
        if signals["slaBreached"] and not data.get("slaAlertedAt"):
            db.add(
                Notification(
                    company_id=lead.organization_id,
                    user_id=lead.owner_user_id or user.id,
                    module="crm",
                    title="Lead response overdue",
                    message=f"Contact {lead.full_name}; the first-response deadline has passed.",
                    event_type="lead_sla",
                    action_url=f"/crm/leads?lead={lead.id}",
                    priority="high",
                )
            )
            data["slaAlertedAt"] = now().isoformat()
            row.data_json = data
            event(db, lead, user, "sla", "First-response deadline escalated")
            count += 1
    db.commit()
    return {"notificationsCreated": count}


def due_follow_up_reminders(db, org, limit):
    """Use the existing protected scheduler; one in-app reminder per pending due task."""
    rows = (
        db.query(CRMActivity, CRMLead)
        .join(CRMLead, CRMLead.id == CRMActivity.lead_id)
        .filter(
            CRMActivity.organization_id == org,
            CRMLead.organization_id == org,
            CRMLead.deleted_at.is_(None),
            CRMLead.is_converted.is_not(True),
            CRMLead.status.notin_(["Lost", "Converted"]),
            CRMActivity.status == "Pending",
            CRMActivity.due_date <= now(),
            CRMActivity.metadata_json["followUpReminderAt"].as_string().is_(None),
        )
        .order_by(CRMActivity.due_date, CRMActivity.id)
        .limit(limit)
        .with_for_update(of=CRMActivity, skip_locked=True)
        .all()
    )
    sent = 0
    from app.apps.crm.lead_assignment import validate_assignee

    for item, lead in rows:
        data = state(db, lead, False)
        if data and data.data_json.get("archived"):
            continue
        recipient = db.get(User, item.owner_user_id) if item.owner_user_id else None
        try:
            validate_assignee(db, recipient, org)
            if not visible_query(db, recipient).filter(CRMLead.id == lead.id).first():
                continue
        except HTTPException:
            continue
        db.add(
            Notification(
                company_id=org,
                user_id=recipient.id,
                module="crm",
                event_type="lead_follow_up_due",
                title="Lead follow-up is due",
                message=f"{lead.full_name}: {item.subject}",
                related_entity_type="lead",
                related_entity_id=lead.id,
                action_url=f"/crm/leads?lead={lead.id}",
                priority="high",
            )
        )
        item.metadata_json = {
            **(item.metadata_json or {}),
            "followUpReminderAt": now().isoformat(),
        }
        sent += 1
    return sent


def scheduled_lead_alerts(db, per_tenant=200):
    """Bounded, tenant-owned SLA escalation for the existing protected cron entry point."""
    organizations = [
        value
        for (value,) in db.query(CRMLead.organization_id)
        .filter(CRMLead.deleted_at.is_(None), CRMLead.organization_id.is_not(None))
        .distinct()
        .all()
    ]
    notifications = 0
    leads_escalated = 0
    follow_up_reminders = 0
    for org in organizations:
        follow_up_reminders += due_follow_up_reminders(db, org, per_tenant)
        settings = config(db, org)
        due = now() - timedelta(hours=settings["slaHours"])
        rows = (
            db.query(CRMLead)
            .outerjoin(
                CRMLeadOperatingState,
                and_(
                    CRMLeadOperatingState.lead_id == CRMLead.id,
                    CRMLeadOperatingState.organization_id == org,
                ),
            )
            .filter(
                CRMLead.organization_id == org,
                CRMLead.deleted_at.is_(None),
                CRMLead.last_contacted_at.is_(None),
                CRMLead.is_converted.is_not(True),
                CRMLead.status.notin_(["Lost", "Converted"]),
                CRMLead.created_at < due,
                CRMLeadOperatingState.data_json["slaAlertedAt"].as_string().is_(None),
                func.coalesce(
                    CRMLeadOperatingState.data_json["archived"].as_boolean(), False
                ).is_(False),
            )
            .order_by(CRMLead.created_at, CRMLead.id)
            .limit(per_tenant)
            .with_for_update(of=CRMLead, skip_locked=True)
            .all()
        )
        setting_row = db.get(CRMLeadManagementConfig, org)
        for lead in rows:
            row = state(db, lead)
            if row.data_json.get("slaAlertedAt"):
                continue
            recipient_ids = {
                lead.owner_user_id,
                setting_row.updated_by if setting_row else None,
            } - {None}
            recipients = [
                candidate
                for candidate in db.query(User)
                .filter(User.id.in_(recipient_ids), User.is_active.is_(True))
                .all()
                if core()._organization_id(db, candidate) == org
            ]
            if not recipients:
                continue
            for recipient in recipients:
                db.add(
                    Notification(
                        company_id=org,
                        user_id=recipient.id,
                        module="crm",
                        title="Lead response overdue",
                        message=f"Contact {lead.full_name}; the first-response deadline has passed.",
                        event_type="lead_sla",
                        action_url=f"/crm/leads?lead={lead.id}",
                        priority="high",
                    )
                )
                notifications += 1
            row.data_json = {**row.data_json, "slaAlertedAt": now().isoformat()}
            row.version += 1
            event(
                db,
                lead,
                recipients[0],
                "sla",
                "First-response deadline escalated automatically",
            )
            leads_escalated += 1
    db.commit()
    return {
        "leadsEscalated": leads_escalated,
        "notificationsCreated": notifications + follow_up_reminders,
        "followUpReminders": follow_up_reminders,
    }


@router.get("/analytics")
def analytics(
    db: Session = Depends(get_db), user: User = Depends(RequirePermission(*ACCESS))
):
    org = core()._organization_id(db, user)
    matching_total=visible_query(db,user).count()
    leads = (
        visible_query(db, user).order_by(CRMLead.id.desc()).limit(MAX_ROWS + 1).all()
    )
    truncated = len(leads) > MAX_ROWS
    leads = leads[:MAX_ROWS]
    ids = [l.id for l in leads]
    states = (
        {
            row.lead_id: row.data_json
            for row in db.query(CRMLeadOperatingState)
            .filter(
                CRMLeadOperatingState.organization_id == org,
                CRMLeadOperatingState.lead_id.in_(ids),
            )
            .all()
        }
        if ids
        else {}
    )
    events = (
        db.query(CRMTimelineEvent)
        .filter(
            CRMTimelineEvent.organization_id == org,
            CRMTimelineEvent.record_type == "lead",
            CRMTimelineEvent.record_id.in_(ids),
            CRMTimelineEvent.event_type.in_(
                ["stage_changed", "conversion", "qualification"]
            ),
        )
        .all()
        if ids
        else []
    )
    durations, entries = {}, {}
    for lead in leads:
        entries.setdefault(lead.status, set()).add(lead.id)
        entries.setdefault("New", set()).add(lead.id)
    for event_row in events:
        metadata = event_row.metadata_json or {}
        previous = metadata.get("from") or metadata.get("previousStatus")
        target = metadata.get("to") or metadata.get("status")
        if previous:
            entries.setdefault(previous, set()).add(event_row.record_id)
        if target:
            entries.setdefault(target, set()).add(event_row.record_id)
        if previous and metadata.get("durationHours") is not None:
            durations.setdefault(previous, []).append(float(metadata["durationHours"]))
    sources, owners = {}, {}
    for lead in leads:
        first = dt(states.get(lead.id, {}).get("firstContactAt"))
        response_minutes = (
            max(0, (first - dt(lead.created_at)).total_seconds() / 60)
            if first and lead.created_at
            else None
        )
        for collection, key in [
            (sources, lead.source or "Other"),
            (owners, str(lead.owner_user_id or "unassigned")),
        ]:
            item = collection.setdefault(
                key,
                {
                    "key": key,
                    "leads": 0,
                    "converted": 0,
                    "qualified": 0,
                    "responseMinutes": [],
                },
            )
            item["leads"] += 1
            item["converted"] += int(bool(lead.is_converted))
            item["qualified"] += int(
                lead.id in entries.get("Qualified", set()) or bool(lead.is_converted)
            )
            if response_minutes is not None:
                item["responseMinutes"].append(response_minutes)
    for collection in [sources, owners]:
        for item in collection.values():
            values = item.pop("responseMinutes")
            item["averageFirstResponseMinutes"] = (
                round(sum(values) / len(values), 1) if values else None
            )
            item["responseSampleSize"] = len(values)
            item["conversionRate"] = round(item["converted"] / item["leads"] * 100, 1)
    return {
        "scope": "All permitted leads; independent of list filters",
        "truncated": truncated,
        "leadCount": len(leads),
        "matchingTotal":matching_total,
        "definition":"Lead performance sample from newest permitted leads; sample aggregates are not whole-directory totals",
        "bySource": list(sources.values()),
        "byOwner": list(owners.values()),
        "stageHistory": [
            {
                "stage": stage["name"],
                "leadsEverInStage": len(entries.get(stage["name"], set())),
                "averageHoursBeforeExit": (
                    round(
                        sum(durations[stage["name"]]) / len(durations[stage["name"]]), 1
                    )
                    if durations.get(stage["name"])
                    else None
                ),
                "timedExits": len(durations.get(stage["name"], [])),
            }
            for stage in config(db, org)["stages"]
        ],
        "measurementNote": "Response averages use first interactions recorded in this workbench. Stage durations use recorded transitions; legacy history may be incomplete.",
    }


def transfer_rows(content, filename, mapping=None, headers_only=False):
    extension = filename.lower().rsplit(".", 1)[-1]
    if extension == "csv":
        try:
            rows = []
            for row in csv.reader(io.StringIO(content.decode("utf-8-sig"))):
                if len(rows) > MAX_ROWS or len(row) > 40:
                    raise HTTPException(
                        422, "Import supports up to 5,000 rows and 40 columns"
                    )
                rows.append(row)
        except (UnicodeError, csv.Error):
            raise HTTPException(422, "Use a UTF-8 CSV or an Excel .xlsx workbook")
    elif extension == "xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if (
                    sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024
                    or len(archive.infolist()) > 1000
                ):
                    raise HTTPException(
                        413, "Workbook is too large after decompression"
                    )
            from openpyxl import load_workbook

            book = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            sheet = book.active
            if sheet.max_column and sheet.max_column > 40:
                book.close()
                raise HTTPException(422, "Import supports up to 40 columns")
            rows = []
            for row in sheet.iter_rows():
                if len(rows) > MAX_ROWS or len(row) > 40:
                    raise HTTPException(
                        422, "Import supports up to 5,000 rows and 40 columns"
                    )
                if any(cell.data_type == "f" for cell in row):
                    raise HTTPException(
                        422, "Replace spreadsheet formulas with values before importing"
                    )
                rows.append([cell.value for cell in row])
            book.close()
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(
                422,
                "This .xlsx workbook could not be read; resave it in Excel and try again",
            )
    else:
        raise HTTPException(
            422,
            "Choose an Excel .xlsx or UTF-8 .csv file. Save legacy .xls files as .xlsx first",
        )
    if len(rows) < 2 or len(rows) > MAX_ROWS + 1:
        raise HTTPException(422, "Include a header and between 1 and 5,000 lead rows")
    aliases = {
        "name": "full_name",
        "lead_name": "full_name",
        "company": "company_name",
        "mobile": "phone",
        "value": "estimated_value",
        "next_follow_up": "next_follow_up_at",
    }
    headers = [
        aliases.get(
            re.sub(r"_+","_",core()._snake(str(h or "").strip().replace(" ", "_"))).strip("_"),
            re.sub(r"_+","_",core()._snake(str(h or "").strip().replace(" ", "_"))).strip("_"),
        )
        for h in rows[0]
    ]
    if headers_only:
        return headers
    if mapping:
        headers=[mapping.get(h,h) for h in headers]
    # Ignore trailing empty cells but reject ambiguous or unsupported data columns.
    while headers and not headers[-1]:
        headers.pop()
    if (
        not headers
        or len([h for h in headers if h]) != len(set(h for h in headers if h))
        or set(h for h in headers if h) - set(TRANSFER_FIELDS)
    ):
        raise HTTPException(
            422,
            {
                "message": "Use unique supported column headers",
                "supportedFields": TRANSFER_FIELDS,
            },
        )
    records = []
    for row in rows[1:]:
        if not any(value is not None and value != "" for value in row):
            continue
        values = {
            key: value.isoformat() if isinstance(value, datetime) else value
            for key, value in zip(headers, row)
            if key and value is not None and value != ""
        }
        # Excel commonly stores national phone numbers as numeric cells. Validate their
        # exact decimal representation; a lost leading zero must be corrected by the user.
        if isinstance(values.get("phone"), (int, float)) and not isinstance(
            values.get("phone"), bool
        ):
            phone = values["phone"]
            values["phone"] = str(int(phone)) if int(phone) == phone else str(phone)
        records.append(values)
    return records


class ImportPayload(Strict):
    rows: list[dict[str, Any]] = Field(min_length=1, max_length=MAX_ROWS)
    skipDuplicates: bool = True


def validate_import(db, user, rows, settings, skip_duplicates=True):
    org = core()._organization_id(db, user)
    clean, errors, skipped = [], [], []
    seen = set()
    for number, row in enumerate(rows, 2):
        try:
            # Excel transfer never accepts owner, audit, tenant or system fields.
            if set(row) - set(TRANSFER_FIELDS):
                raise HTTPException(422, "Unsupported transfer column")
            values = checked_values(
                {
                    key: value
                    for key, value in row.items()
                    if key in BASE_TRANSFER_FIELDS
                }
            )
            qualification_keys = {
                "need": "need",
                "budget": "budget",
                "decision_maker": "decisionMaker",
                "purchase_timeline": "timeline",
            }
            details = {
                "qualification": {
                    target: str(row[key])
                    for key, target in qualification_keys.items()
                    if row.get(key) is not None
                }
            }
            for key, target in {
                "loss_reason": "lossReason",
                "competitor": "competitor",
                "archive_reason": "archiveReason",
                "review_at": "reviewAt",
            }.items():
                if row.get(key) is not None:
                    details[target] = str(row[key])
            if "archived" in row:
                raw_archived = str(row["archived"]).strip().lower()
                if raw_archived not in {"true", "false", "yes", "no", "1", "0"}:
                    raise HTTPException(422, "archived must be true or false")
                details["archived"] = raw_archived in {"true", "yes", "1"}
            details = DetailsPatch(**details).model_dump(
                exclude_unset=True, mode="json"
            )
            validate_transition(settings, None, values, details)
            if details.get("archived") and not details.get("archiveReason"):
                raise HTTPException(422, "Archived leads need an archive_reason")
            keys = (
                {"email:" + str(values["email"]).lower()}
                if values.get("email")
                else set()
            )
            if values.get("phone"):
                keys.add(
                    "phone:"
                    + str(values.get("phone_country_code") or "")
                    + re.sub(r"\D", "", str(values["phone"]))
                )
            duplicate = bool(keys & seen) or bool(duplicates(db, org, values))
            if duplicate and skip_duplicates:
                skipped.append(
                    {
                        "row": number,
                        "reason": "Email or phone already exists in the workbook or CRM",
                    }
                )
                continue
            if duplicate and settings["preventDuplicates"]:
                raise HTTPException(409, "Duplicate lead")
            seen.update(keys)
            clean.append(
                {
                    "row": number,
                    "values": {k: core()._json_ready(v) for k, v in values.items()},
                    "details": details,
                }
            )
        except HTTPException as exc:
            errors.append({"row": number, "error": exc.detail})
        except (ValueError, TypeError, ArithmeticError):
            errors.append(
                {"row": number, "error": "Invalid value; check dates and numbers"}
            )
    return {
        "validRows": clean,
        "errors": errors,
        "skipped": skipped,
        "ready": len(clean),
        "total": len(rows),
    }


@router.post("/excel/columns")
async def import_columns(file:UploadFile=File(...),db:Session=Depends(get_db),user:User=Depends(RequirePermission(*WRITE))):
    core()._assert_entity_permission("leads",user,"manage")
    content=await file.read(MAX_FILE+1)
    if len(content)>MAX_FILE:raise HTTPException(413,"Import file exceeds 3 MB")
    return {"headers":transfer_rows(content,file.filename or "",headers_only=True),"supported_fields":TRANSFER_FIELDS}


@router.post("/excel/preview")
async def import_preview(
    file: UploadFile = File(...),
    field_mapping: str = Form(default=""),
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    core()._assert_entity_permission("leads", user, "manage")
    content = await file.read(MAX_FILE + 1)
    if len(content) > MAX_FILE:
        raise HTTPException(413, "Import file exceeds 3 MB")
    mapping=None
    if field_mapping:
        import json
        try: mapping=json.loads(field_mapping)
        except ValueError: raise HTTPException(422,"Field mapping must be a JSON object")
        if not isinstance(mapping,dict) or not all(isinstance(k,str) and v in [*TRANSFER_FIELDS,""] for k,v in mapping.items()): raise HTTPException(422,"Map each source column to a supported lead field")
    rows = transfer_rows(content, file.filename or "", mapping)
    result = validate_import(
        db, user, rows, config(db, core()._organization_id(db, user))
    )
    result["rows"] = rows
    return result


@router.post("/excel/import")
def import_commit(
    payload: ImportPayload,
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*WRITE)),
):
    settings = config(db, core()._organization_id(db, user), lock=True)
    result = validate_import(db, user, payload.rows, settings, payload.skipDuplicates)
    if result["errors"]:
        raise HTTPException(422, result)
    ids = []
    for row in result["validRows"]:
        # Feed the original supported transfer columns back through validation.
        values = {k: v for k, v in row["values"].items() if k in BASE_TRANSFER_FIELDS}
        lead = create_lead(
            db,
            user,
            LeadPayload(values=values, details=DetailsPatch(**row["details"])),
            settings,
        )
        ids.append(lead.id)
    db.commit()
    return {"created": len(ids), "ids": ids, "skipped": result["skipped"]}


def workbook_response(rows, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "Leads"
    sheet.append(TRANSFER_FIELDS)
    for index, key in enumerate(TRANSFER_FIELDS, 1):
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        sheet.column_dimensions[cell.column_letter].width = 25 if key != "notes" else 50
        if key in {"phone", "phone_country_code", "email"}:
            sheet.column_dimensions[cell.column_letter].number_format = "@"
    for row in rows:
        sheet.append([None] * len(TRANSFER_FIELDS))
        for column, key in enumerate(TRANSFER_FIELDS, 1):
            value = row.get(key)
            cell = sheet.cell(sheet.max_row, column)
            # Explicit string cells prevent formula injection while preserving exact phone/text values.
            if isinstance(value, str):
                value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
                cell.value = value
                cell.data_type = "s"
            else:
                cell.value = value
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    out = io.BytesIO()
    book.save(out)
    book.close()
    return Response(
        out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/excel/template")
def excel_template(user: User = Depends(RequirePermission(*ACCESS))):
    core()._assert_entity_permission("leads", user, "view")
    return workbook_response([], "lead-import-template.xlsx")


@router.get("/excel/export")
def excel_export(
    search: str = "",
    stage: str = "",
    owner: int | None = None,
    ids: str = "",
    db: Session = Depends(get_db),
    user: User = Depends(RequirePermission(*ACCESS)),
):
    query = visible_query(db, user)
    if ids:
        try:
            selected = [int(i) for i in ids.split(",")]
        except ValueError:
            raise HTTPException(422, "Invalid selected lead IDs")
        if len(selected) > 200:
            raise HTTPException(422, "Export at most 200 selected records")
        query = query.filter(CRMLead.id.in_(selected))
    if search:
        query = core()._apply_search(query, core()._get_resource("leads"), search)
    if stage:
        query = query.filter(CRMLead.status == stage)
    if owner is not None:
        query = query.filter(CRMLead.owner_user_id == owner)
    if query.count() > MAX_ROWS:
        raise HTTPException(422, "Narrow your filters to export at most 5,000 leads")
    leads = query.order_by(CRMLead.id).all()
    details_by_lead = (
        {
            s.lead_id: s.data_json
            for s in db.query(CRMLeadOperatingState)
            .filter(
                CRMLeadOperatingState.organization_id
                == core()._organization_id(db, user),
                CRMLeadOperatingState.lead_id.in_([l.id for l in leads]),
            )
            .all()
        }
        if leads
        else {}
    )
    rows = []
    for lead in leads:
        item = core()._serialize(lead)
        details = details_by_lead.get(lead.id, {})
        qualification = details.get("qualification") or {}
        item.update(
            {
                "need": qualification.get("need"),
                "budget": qualification.get("budget"),
                "decision_maker": qualification.get("decisionMaker"),
                "purchase_timeline": qualification.get("timeline"),
                "loss_reason": details.get("lossReason"),
                "competitor": details.get("competitor"),
                "archived": bool(details.get("archived")),
                "archive_reason": details.get("archiveReason"),
                "review_at": details.get("reviewAt"),
            }
        )
        rows.append(item)
    return workbook_response(rows, "leads.xlsx")
